from __future__ import annotations

import json
import zipfile
from pathlib import Path
from typing import Any

import pytest

from review_assist import source_acquisition
from review_assist import deliverable as deliverable_module
from review_assist.data_lineage import build_data_lineage
from review_assist.deliverable import MvpDeliverableError, build_demo_deliverable, build_mvp_deliverable
from review_assist.cli import main
from review_assist.export_report import ExportGateError, ExportQAError, export_report
from review_assist.populate_for_review import populate_for_review
from review_assist.review_queue import load_review_queue, update_review_item


def kml_document(body: str) -> bytes:
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<kml xmlns="http://www.opengis.net/kml/2.2">
  <Document>
    {body}
  </Document>
</kml>
""".encode("utf-8")


def kmz_bytes(kml: bytes) -> bytes:
    import io

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("doc.kml", kml)
    return buffer.getvalue()


def write_project(tmp_path: Path) -> Path:
    project_dir = tmp_path / "project"
    (project_dir / "config").mkdir(parents=True)
    (project_dir / "inputs").mkdir()
    kml = kml_document(
        """
        <Placemark><name>Route A</name><LineString><coordinates>-90.0000,32.0000,0 -89.9900,32.0000,0</coordinates></LineString></Placemark>
        """
    )
    (project_dir / "inputs" / "routes.kmz").write_bytes(kmz_bytes(kml))
    (project_dir / "config" / "project.json").write_text(
        json.dumps(
            {
                "project_id": "test_project",
                "name": "Test Project",
                "description": "Synthetic project",
                "project_type": "alternatives_review",
                "inputs": [
                    {
                        "path": "inputs/routes.kmz",
                        "role": "alternatives",
                        "description": "Synthetic route input",
                    }
                ],
                "assumptions": {
                    "default_buffer_feet": 100,
                    "input_crs": "EPSG:4326",
                },
                "special_reviewer_instructions": "",
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    return project_dir


def fake_nwi_fetch(url: str, params: dict[str, Any]) -> dict[str, Any]:
    if url.endswith("/0"):
        return {"maxRecordCount": 100}
    return {
        "type": "FeatureCollection",
        "features": [
            {
                "type": "Feature",
                "properties": {"name": "Mock NWI Wetland"},
                "geometry": {
                    "type": "Polygon",
                    "coordinates": [
                        [
                            [-90.001, 31.999],
                            [-89.998, 31.999],
                            [-89.998, 32.001],
                            [-90.001, 32.001],
                            [-90.001, 31.999],
                        ]
                    ],
                },
            }
        ],
    }


def fake_empty_public_fetch(url: str, params: dict[str, Any]) -> dict[str, Any]:
    if params.get("f") == "pjson":
        return {"maxRecordCount": 100}
    return {"type": "FeatureCollection", "features": []}


def add_source_input(project_dir: Path, source_id: str, filename: str, geometry: dict[str, Any]) -> None:
    input_path = project_dir / "inputs" / filename
    input_path.write_text(
        json.dumps(
            {
                "type": "FeatureCollection",
                "features": [
                    {
                        "type": "Feature",
                        "properties": {"name": source_id},
                        "geometry": geometry,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    manifest_path = project_dir / "config" / "project.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["inputs"].append(
        {
            "path": f"inputs/{filename}",
            "role": "source_layer",
            "description": f"Provided source layer for {source_id}",
            "source_id": source_id,
        }
    )
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")


def add_supported_real_source_inputs(project_dir: Path) -> None:
    add_source_input(
        project_dir,
        "usfws_nwi_wetlands",
        "provided_nwi.geojson",
        {
            "type": "Polygon",
            "coordinates": [[[-90.001, 31.999], [-89.998, 31.999], [-89.998, 32.001], [-90.001, 32.001], [-90.001, 31.999]]],
        },
    )
    add_source_input(
        project_dir,
        "usgs_nhd_hydrography",
        "provided_nhd.geojson",
        {
            "type": "LineString",
            "coordinates": [[-90.001, 32.0], [-89.998, 32.0]],
        },
    )
    add_source_input(
        project_dir,
        "usfws_critical_habitat",
        "provided_critical_habitat.geojson",
        {
            "type": "Polygon",
            "coordinates": [[[-90.002, 31.998], [-89.997, 31.998], [-89.997, 32.002], [-90.002, 32.002], [-90.002, 31.998]]],
        },
    )
    add_source_input(
        project_dir,
        "epa_envirofacts_echo",
        "provided_echo.geojson",
        {
            "type": "Point",
            "coordinates": [-89.999, 32.0],
        },
    )


def set_queue_item(project_dir: Path, item_id: str, **updates: Any) -> None:
    queue = load_review_queue(project_dir)
    for item in queue["items"]:
        if item["id"] == item_id or item.get("target_id") == item_id or item.get("deliverable_item_id") == item_id:
            item.update(updates)
            Path(queue["output_path"]).write_text(json.dumps(queue, indent=2) + "\n", encoding="utf-8")
            return
    raise AssertionError(f"Missing queue item: {item_id}")


def included_ids(manifest: dict[str, Any]) -> set[str]:
    return {item["id"] for item in manifest["included_items"]}


def qa_codes(error: ExportQAError) -> set[str]:
    return {str(issue.get("code")) for issue in error.details.get("export_qa_issues", [])}


def mutate_queue_item(project_dir: Path, item_id: str, callback: Any) -> None:
    queue = load_review_queue(project_dir)
    for item in queue["items"]:
        if item["id"] == item_id or item.get("target_id") == item_id or item.get("deliverable_item_id") == item_id:
            callback(item)
            Path(queue["output_path"]).write_text(json.dumps(queue, indent=2) + "\n", encoding="utf-8")
            return
    raise AssertionError(f"Missing queue item: {item_id}")


def docx_text(path: str | Path) -> str:
    from docx import Document

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def docx_document(path: str | Path) -> Any:
    from docx import Document

    return Document(path)


def write_tiny_png(path: Path) -> None:
    import base64

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAAEklEQVR4nGNgYGD4z8DAwMAAAAUgAgnqQZ7nAAAAAElFTkSuQmCC"
        )
    )


def set_review_states(
    project_dir: Path,
    *,
    default_status: str = "accepted",
    overrides: dict[str, str | dict[str, Any]] | None = None,
) -> None:
    overrides = overrides or {}
    queue = load_review_queue(project_dir)
    for item in queue["items"]:
        spec = overrides.get(item["id"]) or overrides.get(item.get("target_id", "")) or overrides.get(item.get("deliverable_item_id", ""))
        if isinstance(spec, str):
            spec = {"status": spec}
        elif spec is None:
            spec = {}
        status = str(spec.get("status", default_status))
        item["status"] = status
        if "generated_content" in spec:
            item["generated_content"] = spec["generated_content"]
        if "edited_content" in spec:
            item["edited_content"] = spec["edited_content"]
        if "replacement_content" in spec:
            item["replacement_content"] = spec["replacement_content"]
        if "export_eligible" in spec:
            item["export_eligible"] = bool(spec["export_eligible"])
        elif status in {"accepted", "edited"}:
            item["export_eligible"] = True
        elif status == "replaced":
            item["export_eligible"] = bool(str(item.get("replacement_content", "")).strip())
        elif status == "unable_to_verify":
            item["export_eligible"] = False
        else:
            item["export_eligible"] = False
    Path(queue["output_path"]).write_text(json.dumps(queue, indent=2) + "\n", encoding="utf-8")


def set_only_reviewed_items(project_dir: Path, overrides: dict[str, str | dict[str, Any]]) -> None:
    set_review_states(project_dir, default_status="declined", overrides=overrides)


def test_docx_export_uses_letter_page_setup_and_core_styles(tmp_path: Path) -> None:
    from docx.shared import Inches, Pt, RGBColor

    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir)

    manifest = export_report(project_dir, output_format="docx")
    document = docx_document(manifest["docx_path"])
    section = document.sections[0]

    assert section.page_width == Inches(8.5)
    assert section.page_height == Inches(11)
    assert section.top_margin == Inches(1)
    assert section.bottom_margin == Inches(1)
    assert section.left_margin == Inches(1)
    assert section.right_margin == Inches(1)
    assert section.header_distance == Inches(0.5)
    assert section.footer_distance == Inches(0.5)

    styles = document.styles
    assert styles["Normal"].font.name == "Calibri"
    assert styles["Normal"].font.size == Pt(12)
    assert styles["Normal"].paragraph_format.space_after == Pt(8)
    assert styles["Heading 1"].font.name == "Lato"
    assert styles["Heading 1"].font.size == Pt(20)
    assert styles["Heading 1"].font.color.rgb == RGBColor(0x0F, 0x47, 0x61)
    assert styles["Heading 2"].font.size == Pt(14)
    assert styles["Caption"].font.name == "Calibri Light"
    assert styles["Caption"].font.size == Pt(11)
    assert styles["Attachment Title"].font.color.rgb == RGBColor(0x0F, 0x47, 0x61)


def test_docx_export_title_front_matter_and_outline_follow_matrix(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir)

    manifest = export_report(project_dir, output_format="docx")
    text = docx_text(manifest["docx_path"])
    queue_ids = {item["id"] for item in load_review_queue(project_dir)["items"]}
    natural_heading = next(item for item in manifest["included_items"] if item["id"] == "natural-and-ecological-resources")

    assert "Test Project" in text
    assert "Project ID: test_project" in text
    assert "Synthetic project" in text
    assert "INTERNAL PREVIEW / NOT REVIEWED" not in text
    assert "List of Figures" in text
    assert "Wetlands and Waterbodies in and near the Project Area" in text
    assert "List of Tables" in text
    assert "Descriptions of Wetlands and Waterbodies Present within the Project Area" in text
    assert "List of Attachments" in text
    assert text.index("Attachment A") < text.index("Attachment B") < text.index("Attachment C")
    assert "natural-and-ecological-resources" not in queue_ids
    assert natural_heading["type"] == "structural_heading"
    assert natural_heading["status"] == "not_review_required"
    assert natural_heading["content"] == ""
    assert text.index("3.1 Natural and Ecological Resources") < text.index("3.1.1 Wetlands and Waterbodies")

    expected_outline = [
        "Executive Summary",
        "Introduction",
        "Methodology",
        "Environmental Constraints Inventory",
        "3.1 Natural and Ecological Resources",
        "Community Resources",
        "Utility and Infrastructure Considerations",
        "Contamination Risks",
        "Socioeconomic and Business Considerations",
        "Conclusion and Next Steps",
        "Attachments",
    ]
    document = docx_document(manifest["docx_path"])
    headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style and paragraph.style.name.startswith("Heading")]
    positions = [headings.index(title) for title in expected_outline]
    assert positions == sorted(positions)
    skipped_cultural = next(item for item in manifest["skipped_items"] if item["id"] == "cultural-and-historic-resources")
    assert skipped_cultural["reason"] == "policy_render_blocked_manual_or_restricted_source"


def test_docx_preview_label_appears_only_in_preview_mode(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    preview_manifest = export_report(project_dir, include_draft=True, output_format="docx")
    preview_text = docx_text(preview_manifest["docx_path"])
    assert "INTERNAL PREVIEW / NOT REVIEWED" in preview_text

    set_review_states(project_dir)
    reviewed_manifest = export_report(project_dir, output_format="docx")
    reviewed_text = docx_text(reviewed_manifest["docx_path"])
    assert "INTERNAL PREVIEW / NOT REVIEWED" not in reviewed_text


def test_docx_renders_reviewed_table_preview_and_figure_metadata_compactly(tmp_path: Path) -> None:
    from test_deliverable_compactness import _write_large_deliverable_table

    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    _write_large_deliverable_table(project_dir, row_count=60)
    from review_assist.deliverable_items import generate_deliverable_items
    from review_assist.review_queue import generate_review_queue

    generate_deliverable_items(project_dir, gpt_drafting=False)
    generate_review_queue(project_dir)
    figure_path = project_dir / "maps" / "figures" / "wetlands-test.png"
    write_tiny_png(figure_path)
    set_queue_item(
        project_dir,
        "figure-wetlands-waterbodies",
        image_path=str(figure_path),
        caption="Wetlands figure caption.",
        source_note="Wetlands source note.",
        method_note="Wetlands method note.",
    )
    set_only_reviewed_items(
        project_dir,
        {
            "wetlands-and-waterbodies": "accepted",
            "table-wetlands-waterbodies": "accepted",
            "figure-wetlands-waterbodies": "accepted",
        },
    )

    manifest = export_report(project_dir, output_format="docx")
    text = docx_text(manifest["docx_path"])

    assert "Table preview limited to 5 of 60 rows" in text
    assert "Alternative 004" in text
    assert "Alternative 059" not in text
    assert "Wetlands figure caption." in text
    assert "Wetlands source note." in text
    assert "Wetlands method note." in text
    assert manifest["final_verification"]["docx_readable"] is True
    assert manifest["final_verification"]["raw_legacy_item_count"] == 0
    assert manifest["final_verification"]["max_table_preview_rows"] <= 5


def test_final_verification_flags_raw_legacy_and_over_budget_content(tmp_path: Path) -> None:
    from review_assist.review_queue import generate_review_queue

    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    generate_review_queue(project_dir, include_legacy_artifacts=True)
    legacy_manifest = export_report(project_dir, include_draft=True)
    assert legacy_manifest["final_verification"]["status"] == "failed"
    assert legacy_manifest["final_verification"]["raw_legacy_item_count"] > 0
    assert any(issue["code"] == "raw_legacy_items_in_export" for issue in legacy_manifest["final_verification"]["issues"])

    generate_review_queue(project_dir)
    set_queue_item(project_dir, "study-area", generated_content="Oversized reviewed section. " * 250)
    set_review_states(project_dir)
    reviewed_manifest = export_report(project_dir)

    assert any(issue["code"] == "export_body_content_over_budget" for issue in reviewed_manifest["final_verification"]["issues"])


def test_sprint_3_3_smoke_gate_then_terminal_docx_export_succeeds(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    with pytest.raises(ExportGateError):
        export_report(project_dir, output_format="docx")

    set_review_states(project_dir)
    manifest = export_report(project_dir, output_format="both")

    assert manifest["review_gate_status"] == "passed"
    assert manifest["export_qa_blocking_error_count"] == 0
    assert manifest["final_verification"]["status"] == "passed"
    assert Path(manifest["markdown_path"]).exists()
    assert Path(manifest["docx_path"]).exists()
    assert manifest["final_verification"]["docx_readable"] is True


def test_export_manifest_filters_reviewed_items_and_uses_edited_content(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    set_review_states(
        project_dir,
        overrides={
            "executive-summary": {"status": "edited", "edited_content": "Reviewer edited executive summary."},
            "introduction": "declined",
        },
    )

    manifest = export_report(project_dir)
    markdown = Path(manifest["markdown_path"]).read_text(encoding="utf-8")

    assert "cover-title" in included_ids(manifest)
    assert "executive-summary" in included_ids(manifest)
    assert "introduction" not in included_ids(manifest)
    assert "study-area" in included_ids(manifest)
    assert "Reviewer edited executive summary." in markdown
    assert "INTERNAL PREVIEW EXPORT" not in markdown
    assert manifest["review_gate_status"] == "passed"
    assert manifest["preview_mode"] is False
    assert manifest["unreviewed_item_count"] == 0
    assert manifest["deliverable_matrix_version"]
    assert manifest["expected_deliverable_item_count"] == manifest["actual_deliverable_item_count"]
    assert manifest["included_table_ids"]
    assert manifest["included_figure_ids"]
    assert manifest["included_attachment_ids"]
    assert manifest["stub_item_count"] > 0


def test_preview_export_records_export_qa_without_hard_blocking(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    manifest = export_report(project_dir, include_draft=True)

    assert manifest["preview_mode"] is True
    assert manifest["export_qa_status"] in {"warning", "failed"}
    assert manifest["export_qa"]["override_supported"] is False
    assert any(issue["code"] == "preview_export_qa_bypassed" for issue in manifest["export_qa_issues"])


def test_reviewed_export_hard_blocks_unknown_source_ref(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir)
    set_queue_item(project_dir, "study-area", source_refs=["not_a_source_id"])

    with pytest.raises(ExportQAError) as exc:
        export_report(project_dir)

    assert "unknown_source_ref" in qa_codes(exc.value)
    assert not (project_dir / "exports" / "environmental_constraints_report.md").exists()


def test_reviewed_export_hard_blocks_disallowed_figure_ref(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir)
    set_queue_item(project_dir, "wetlands-and-waterbodies", related_figure_ids=["figure-fema-flood-zones"])

    with pytest.raises(ExportQAError) as exc:
        export_report(project_dir)

    assert "disallowed_figure_ref" in qa_codes(exc.value)


def test_reviewed_export_hard_blocks_missing_required_caveat_metadata(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir, default_status="declined", overrides={"wetlands-and-waterbodies": "accepted"})

    def remove_caveats(item: dict[str, Any]) -> None:
        item["required_caveats"] = []
        assumptions = item.get("assumptions", {}) if isinstance(item.get("assumptions"), dict) else {}
        assumptions["required_caveats"] = []
        matrix_target = assumptions.get("matrix_target", {}) if isinstance(assumptions.get("matrix_target"), dict) else {}
        matrix_target["required_caveats"] = []
        assumptions["matrix_target"] = matrix_target
        item["assumptions"] = assumptions

    mutate_queue_item(project_dir, "wetlands-and-waterbodies", remove_caveats)

    with pytest.raises(ExportQAError) as exc:
        export_report(project_dir)

    assert "required_caveat_missing" in qa_codes(exc.value)


def test_reviewed_export_hard_blocks_blank_required_table_cell(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir, default_status="declined", overrides={"table-wetlands-waterbodies": "accepted"})
    set_queue_item(
        project_dir,
        "table-wetlands-waterbodies",
        is_stub=False,
        table_id="table-wetlands-waterbodies",
        columns=["Feature", "Result"],
        rows_preview=[{"Feature": "", "Result": "reviewed"}],
        required_columns=["Feature"],
    )

    with pytest.raises(ExportQAError) as exc:
        export_report(project_dir)

    assert "blank_required_table_cell" in qa_codes(exc.value)


def test_reviewed_export_hard_blocks_missing_figure_metadata(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir, default_status="declined", overrides={"figure-wetlands-waterbodies": "accepted"})

    def remove_figure_metadata(item: dict[str, Any]) -> None:
        item.update({"is_stub": False, "caption": "", "source_note": "", "method_note": "", "image_path": ""})
        assumptions = item.get("assumptions", {}) if isinstance(item.get("assumptions"), dict) else {}
        assumptions.update({"caption": "", "source_note": "", "method_note": "", "image_path": ""})
        item["assumptions"] = assumptions

    mutate_queue_item(project_dir, "figure-wetlands-waterbodies", remove_figure_metadata)

    with pytest.raises(ExportQAError) as exc:
        export_report(project_dir)

    assert {"missing_figure_caption", "missing_figure_source_note", "missing_figure_method_note"}.issubset(qa_codes(exc.value))


def test_reviewed_export_hard_blocks_manual_material_without_reviewer_content(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir, default_status="declined", overrides={"relationship-with-pel-study": "accepted"})
    set_queue_item(
        project_dir,
        "relationship-with-pel-study",
        report_body_eligible=True,
        render_decision="include_body",
        render_destination="report_body",
        export_eligible=True,
        manual_material={
            "material_type": "manual_text",
            "material_status": "manual_required",
            "export_behavior": "body_replacement_when_reviewed",
            "reviewer_action": "Supply reviewer-approved parent-study relationship text before export.",
            "source_refs": [],
            "source_categories": [],
            "internal_note_only": False,
        },
    )

    with pytest.raises(ExportQAError) as exc:
        export_report(project_dir)

    assert "manual_material_without_reviewer_content" in qa_codes(exc.value)


def test_export_edited_without_content_falls_back_with_warning(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir, overrides={"study-area": {"status": "edited", "edited_content": ""}})

    manifest = export_report(project_dir)
    exported = next(item for item in manifest["included_items"] if item["id"] == "study-area")

    assert exported["content_source"] == "generated_content"
    assert any(issue["code"] == "edited_content_missing" and issue["item_id"] == "study-area" for issue in manifest["validation_issues"])


def test_export_replaced_requires_replacement_content_and_exports_replacement(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir, overrides={"study-area": {"status": "replaced", "replacement_content": ""}})

    with pytest.raises(ExportGateError) as exc:
        export_report(project_dir)
    assert any(item["reason"] == "replacement_content_missing" for item in exc.value.details["unreviewed_items_preview"])

    set_review_states(
        project_dir,
        overrides={"study-area": {"status": "replaced", "replacement_content": "Reviewer replacement study area."}},
    )
    manifest = export_report(project_dir)
    exported = next(item for item in manifest["included_items"] if item["id"] == "study-area")

    assert exported["content"] == "Reviewer replacement study area."
    assert exported["content_source"] == "replacement_content"


def test_export_figure_uses_edited_caption_with_generated_image(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    image_path = project_dir / "maps" / "figures" / "generated-figure.png"
    write_tiny_png(image_path)
    set_review_states(
        project_dir,
        default_status="declined",
        overrides={
            "figure-wetlands-waterbodies": {
                "status": "edited",
                "edited_content": "Reviewed wetlands and waterbodies caption.",
                "export_eligible": True,
            }
        },
    )
    set_queue_item(project_dir, "figure-wetlands-waterbodies", image_path="maps/figures/generated-figure.png")

    manifest = export_report(project_dir)
    markdown = Path(manifest["markdown_path"]).read_text(encoding="utf-8")
    exported = next(item for item in manifest["included_items"] if item["id"] == "figure-wetlands-waterbodies")

    assert exported["caption"] == "Reviewed wetlands and waterbodies caption."
    assert exported["caption_source"] == "edited_caption"
    assert exported["image_source"] == "generated_figure"
    assert exported["content"] == ""
    assert "Caption: Reviewed wetlands and waterbodies caption." in markdown
    assert "reviewer verification" not in markdown.lower()
    assert "draft figure review" not in markdown.lower()


def test_export_figure_uses_replacement_image_and_edited_caption(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    replacement_path = project_dir / "review_queue" / "figure_replacements" / "figure-wetlands-waterbodies" / "replacement.png"
    write_tiny_png(replacement_path)
    replacement_rel = replacement_path.relative_to(project_dir).as_posix()
    set_review_states(
        project_dir,
        default_status="declined",
        overrides={
            "figure-wetlands-waterbodies": {
                "status": "replaced",
                "edited_content": "Replacement wetlands figure caption.",
                "replacement_content": replacement_rel,
                "export_eligible": True,
            }
        },
    )
    set_queue_item(project_dir, "figure-wetlands-waterbodies", reviewer_notes=[{"created_at": "now", "note": "Internal reviewer note only."}])

    manifest = export_report(project_dir)
    markdown = Path(manifest["markdown_path"]).read_text(encoding="utf-8")
    exported = next(item for item in manifest["included_items"] if item["id"] == "figure-wetlands-waterbodies")
    asset = next(item for item in manifest["export_figure_assets"] if item["figure_id"] == "figure-wetlands-waterbodies")

    assert exported["caption"] == "Replacement wetlands figure caption."
    assert exported["image_path"] == replacement_rel
    assert exported["caption_source"] == "edited_caption"
    assert exported["image_source"] == "replacement_figure"
    assert exported["content_source"] == "replacement_figure"
    assert Path(asset["source_image_path"]) == replacement_path
    assert asset["export_asset_path"] == "assets/figures/figure-wetlands-waterbodies.png"
    assert "Caption: Replacement wetlands figure caption." in markdown
    assert "Map file: `assets/figures/figure-wetlands-waterbodies.png`" in markdown
    assert "Internal reviewer note only" not in markdown


def test_export_figure_replacement_uses_generated_caption_when_caption_is_not_edited(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    queue = load_review_queue(project_dir)
    existing = next(item for item in queue["items"] if item["id"] == "figure-wetlands-waterbodies")
    generated_caption = existing["assumptions"]["caption"]
    replacement_path = project_dir / "review_queue" / "figure_replacements" / "figure-wetlands-waterbodies" / "replacement.png"
    write_tiny_png(replacement_path)
    replacement_rel = replacement_path.relative_to(project_dir).as_posix()
    set_review_states(
        project_dir,
        default_status="declined",
        overrides={
            "figure-wetlands-waterbodies": {
                "status": "replaced",
                "replacement_content": replacement_rel,
                "export_eligible": True,
            }
        },
    )

    manifest = export_report(project_dir)
    markdown = Path(manifest["markdown_path"]).read_text(encoding="utf-8")
    exported = next(item for item in manifest["included_items"] if item["id"] == "figure-wetlands-waterbodies")

    assert exported["caption"] == generated_caption
    assert exported["caption_source"] == "generated_caption"
    assert exported["image_source"] == "replacement_figure"
    assert exported["content_source"] == "replacement_figure"
    assert f"Caption: {generated_caption}" in markdown


def test_export_legacy_rejected_status_is_omitted_as_declined(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir, overrides={"introduction": "rejected"})

    manifest = export_report(project_dir)

    assert "introduction" not in included_ids(manifest)
    assert manifest["declined_item_count"] == 1
    assert manifest["review_gate_status"] == "passed"


def test_data_lineage_ignores_stale_download_manifest_without_active_registry_source(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    acquisition_dir = project_dir / "source_acquisition"
    acquisition_dir.mkdir()
    (acquisition_dir / "source_acquisition_manifest.json").write_text(
        json.dumps(
            {
                "downloads": [
                    {
                        "source_id": "usgs_nhd_hydrography",
                        "source_name": "National Hydrography Dataset",
                        "source_category": "hydrography_crossings",
                        "status": "downloaded",
                        "data_authenticity": "real",
                        "output_path": str(project_dir / "source_acquisition" / "downloads" / "usgs_nhd_hydrography.geojson"),
                        "feature_count": 12,
                    }
                ]
            }
        ),
        encoding="utf-8",
    )

    lineage = build_data_lineage(project_dir)

    assert lineage["counts"]["downloaded_public_source"] == 0
    assert lineage["real_source_count"] == 0
    assert any(issue["code"] == "stale_download_record_ignored" for issue in lineage["validation_issues"])


def test_export_includes_unable_to_verify_only_when_export_eligible(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    set_review_states(project_dir, overrides={"limitations-and-data-gaps": "unable_to_verify"})
    with pytest.raises(ExportGateError) as exc:
        export_report(project_dir)
    assert exc.value.details["review_gate_status"] == "blocked"
    assert any(item["id"] == "limitations-and-data-gaps" for item in exc.value.details["unreviewed_items_preview"])

    set_review_states(
        project_dir,
        overrides={"limitations-and-data-gaps": {"status": "unable_to_verify", "export_eligible": True}},
    )
    second = export_report(project_dir)
    assert "limitations-and-data-gaps" in included_ids(second)
    assert second["review_gate_status"] == "passed"


def test_export_skips_body_ineligible_generated_placeholders_until_reviewer_supplies_content(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    set_review_states(project_dir)
    manifest = export_report(project_dir)

    assert "relationship-with-pel-study" not in included_ids(manifest)
    skipped_pel = next(item for item in manifest["skipped_items"] if item["id"] == "relationship-with-pel-study")
    assert skipped_pel["render_decision"] == "needs_reviewer_decision"
    assert skipped_pel["report_body_eligible"] is False
    assert skipped_pel["reason"] == "policy_render_needs_reviewer_decision"
    assert skipped_pel["manual_material"]["material_type"] == "manual_text"
    assert skipped_pel["manual_material"]["material_status"] == "manual_required"

    set_review_states(
        project_dir,
        overrides={
            "relationship-with-pel-study": {
                "status": "edited",
                "edited_content": "Reviewer supplied parent-study relationship content.",
                "export_eligible": True,
            }
        },
    )
    reviewed = export_report(project_dir)

    exported_pel = next(item for item in reviewed["included_items"] if item["id"] == "relationship-with-pel-study")
    assert exported_pel["content"] == "Reviewer supplied parent-study relationship content."
    assert exported_pel["content_source"] == "edited_content"
    assert exported_pel["render_decision"] == "needs_reviewer_decision"
    assert exported_pel["report_body_eligible"] is False
    assert exported_pel["manual_material"]["material_status"] == "reviewer_supplied"
    assert exported_pel["manual_material"]["material_type"] == "manual_text"


def test_default_export_fails_when_review_gate_has_unreviewed_items(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    with pytest.raises(ExportGateError) as exc:
        export_report(project_dir)

    details = exc.value.details
    assert details["review_gate_status"] == "blocked"
    assert details["review_item_count"] > 0
    assert details["unreviewed_item_count"] > 0
    assert details["unreviewed_items_preview"]
    assert "include-draft" in str(exc.value)


@pytest.mark.parametrize("status", ["draft", "needs_review", "needs_verification"])
def test_default_export_gate_blocks_each_unreviewed_status(tmp_path: Path, status: str) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir, overrides={"study-area": status})

    with pytest.raises(ExportGateError) as exc:
        export_report(project_dir)

    assert exc.value.details["review_gate_status"] == "blocked"
    assert any(item["id"] == "study-area" and item["status"] == status for item in exc.value.details["unreviewed_items_preview"])


def test_export_preview_includes_drafts_and_marks_markdown(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    manifest = export_report(project_dir, include_draft=True)
    markdown = Path(manifest["markdown_path"]).read_text(encoding="utf-8")

    assert manifest["include_draft"] is True
    assert manifest["preview_mode"] is True
    assert manifest["review_gate_status"] == "preview_bypassed"
    assert manifest["included_count"] > 0
    assert "study-area" in included_ids(manifest)
    assert "INTERNAL PREVIEW EXPORT" in markdown


def test_export_docx_preview_includes_drafts_and_marks_output(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    manifest = export_report(project_dir, include_draft=True, output_format="docx")
    text = docx_text(manifest["docx_path"])

    assert manifest["markdown_path"] is None
    assert manifest["docx_path"].endswith("environmental_constraints_report.docx")
    assert Path(manifest["docx_path"]).exists()
    assert "docx" in manifest["output_formats"]
    assert "INTERNAL PREVIEW / NOT REVIEWED" in text
    assert "Real Data Used" in text
    assert "Stubs / Manual Review Needed" in text
    assert manifest["data_lineage"]["counts"]["project_input"] == 1
    assert manifest["data_lineage"]["counts"]["missing_stub"] > 0
    assert any(issue["code"] == "no_real_source_layers" for issue in manifest["validation_issues"])
    assert "Study Area" in text
    assert "Generated Package Contents" in text


def test_export_strips_duplicate_report_section_headings(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_queue_item(
        project_dir,
        "cover-title",
        generated_content="# Front Matter\n\nBody without duplicate heading.",
    )

    manifest = export_report(project_dir, include_draft=True, output_format="both")
    markdown = Path(manifest["markdown_path"]).read_text(encoding="utf-8")
    text = docx_text(manifest["docx_path"])

    assert "\n### Front Matter\n" not in markdown
    assert "Body without duplicate heading." in markdown
    assert "Front Matter\nFront Matter\nBody without duplicate heading." not in text
    assert "Body without duplicate heading." in text


def test_docx_export_front_matter_lists_included_figures_tables_and_attachments(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    manifest = export_report(project_dir, include_draft=True, output_format="docx")
    text = docx_text(manifest["docx_path"])

    assert "List of Figures" in text
    assert "Wetlands and Waterbodies in and near the Project Area" in text
    assert "List of Tables" in text
    assert "Descriptions of Wetlands and Waterbodies Present within the Project Area" in text
    assert "List of Attachments" in text
    assert "Attachment A: Project Maps." in text


def test_preview_export_front_matter_lists_use_included_table_and_figure_labels(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    manifest = export_report(project_dir, include_draft=True, output_format="markdown")
    markdown = Path(manifest["markdown_path"]).read_text(encoding="utf-8")
    figure_list = markdown.split("### List of Figures", 1)[1].split("### List of Tables", 1)[0]
    table_list = markdown.split("### List of Tables", 1)[1].split("### List of Attachments", 1)[0]

    assert "- Figure 1. Wetlands and Waterbodies in and near the Project Area" in figure_list
    assert "- Figure 2. FEMA Flood Zones in and near the Project Area" in figure_list
    assert "- Table 1. Descriptions of Wetlands and Waterbodies Present within the Project Area" in table_list
    assert "- Table 2. FEMA Flood Zones within the Project Area" in table_list
    assert "The figure list reflects matrix-backed figure items in matrix order." not in figure_list
    assert "The table list reflects matrix-backed table items in matrix order." not in table_list
    assert "`figure-wetlands-waterbodies`" not in figure_list
    assert "`table-wetlands-waterbodies`" not in table_list


def test_reviewed_export_front_matter_lists_only_export_eligible_reviewed_artifacts(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_only_reviewed_items(
        project_dir,
        {
            "list-of-figures": "accepted",
            "list-of-tables": "accepted",
            "figure-streams-impaired-waters": "accepted",
            "table-fema-flood-zones": "accepted",
        },
    )

    manifest = export_report(project_dir, output_format="markdown")
    markdown = Path(manifest["markdown_path"]).read_text(encoding="utf-8")

    assert "- Figure 3. Streams and 303(d) Impaired Waters within Project-Area Subwatersheds" in markdown
    assert "- Figure 1. Wetlands and Waterbodies in and near the Project Area" not in markdown
    assert "- Table 2. FEMA Flood Zones within the Project Area" in markdown
    assert "- Table 1. Descriptions of Wetlands and Waterbodies Present within the Project Area" not in markdown


def test_export_format_both_writes_markdown_and_docx(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    manifest = export_report(project_dir, include_draft=True, output_format="both")

    assert Path(manifest["markdown_path"]).exists()
    assert Path(manifest["docx_path"]).exists()
    assert manifest["output_formats"] == ["markdown", "docx"]


def test_docx_export_embeds_table_content_and_missing_figure_placeholder(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_queue_item(project_dir, "figure-wetlands-waterbodies", image_path=str(project_dir / "missing-map.png"))

    manifest = export_report(project_dir, include_draft=True, output_format="docx")
    text = docx_text(manifest["docx_path"])

    assert "Descriptions of Wetlands and Waterbodies Present within the Project Area" in text
    assert "Table placeholder: this table currently has no rows." in text
    assert "Figure placeholder: figure file was not available" in text
    assert any(issue["code"] == "missing_export_figure_asset" for issue in manifest["validation_issues"])
    assert manifest["mvp_quality"]["missing_figure_asset_warning_count"] >= 1


def test_export_with_nwi_backed_constraints_includes_accepted_findings_tables_and_maps(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    project_dir = write_project(tmp_path)
    monkeypatch.setattr(source_acquisition, "_fetch_json", fake_nwi_fetch)
    populate_for_review(project_dir, prepare_sources=True)
    finding_id = "wetlands-and-waterbodies"
    map_id = "figure-wetlands-waterbodies"

    set_only_reviewed_items(
        project_dir,
        {
            finding_id: "accepted",
            "table-wetlands-waterbodies": "accepted",
            map_id: "accepted",
        },
    )

    manifest = export_report(project_dir)
    markdown = Path(manifest["markdown_path"]).read_text(encoding="utf-8")

    assert finding_id in included_ids(manifest)
    assert "table-wetlands-waterbodies" in included_ids(manifest)
    assert map_id in included_ids(manifest)
    assert manifest["data_lineage"]["counts"]["test_or_mock"] > 0
    assert "table-wetlands-waterbodies" in markdown
    assert "Map file:" in markdown
    assert "Caption:" in markdown
    assert "Source note:" in markdown
    assert manifest["export_figure_assets"]
    assert manifest["mvp_quality"]["copied_figure_asset_count"] == 1
    assert Path(manifest["export_figure_assets"][0]["export_image_path"]).exists()


def test_docx_export_with_nwi_backed_constraints_includes_accepted_evidence(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    project_dir = write_project(tmp_path)
    monkeypatch.setattr(source_acquisition, "_fetch_json", fake_nwi_fetch)
    populate_for_review(project_dir, prepare_sources=True)
    finding_id = "wetlands-and-waterbodies"
    map_id = "figure-wetlands-waterbodies"

    set_only_reviewed_items(
        project_dir,
        {
            finding_id: "accepted",
            "table-wetlands-waterbodies": "accepted",
            map_id: "accepted",
        },
    )

    manifest = export_report(project_dir, output_format="docx")
    text = docx_text(manifest["docx_path"])

    assert finding_id in included_ids(manifest)
    assert "Wetlands and Waterbodies" in text
    assert "Descriptions of Wetlands and Waterbodies" in text
    assert "Figure file:" in text
    assert "Caption:" in text
    assert "Source note:" in text
    assert manifest["mvp_quality"]["inline_rendered_table_count"] > 0
    assert manifest["mvp_quality"]["inline_rendered_figure_count"] > 0
    assert manifest["mvp_quality"]["copied_figure_asset_count"] == 1
    assert text.count("Figure file:") == 1


def test_report_sections_render_related_table_and_figure_inline_without_standalone_duplicates(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    project_dir = write_project(tmp_path)
    monkeypatch.setattr(source_acquisition, "_fetch_json", fake_nwi_fetch)
    populate_for_review(project_dir, prepare_sources=True)
    map_id = "figure-wetlands-waterbodies"

    set_only_reviewed_items(
        project_dir,
        {
            "wetlands-and-waterbodies": "accepted",
            "table-wetlands-waterbodies": "accepted",
            map_id: "accepted",
        },
    )

    manifest = export_report(project_dir, output_format="docx")
    text = docx_text(manifest["docx_path"])

    assert "Table: Descriptions of Wetlands and Waterbodies" in text
    assert "Figure file:" in text
    assert "Generated deliverable figure" not in text
    assert manifest["mvp_quality"]["inline_rendered_table_ids"] == ["table-wetlands-waterbodies"]
    assert manifest["mvp_quality"]["inline_rendered_figure_count"] == 1
    assert manifest["mvp_quality"]["copied_figure_asset_ids"] == ["figure-wetlands-waterbodies"]


def test_export_warns_when_existing_map_manifest_cannot_be_loaded(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    manifest_path = project_dir / "maps" / "map_manifest.json"
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text("{not-json", encoding="utf-8")

    manifest = export_report(project_dir, include_draft=True)

    assert any(issue["code"] == "map_manifest_unavailable" for issue in manifest["validation_issues"])
    assert manifest["mvp_quality"]["validation_warning_count"] >= 1


def test_export_preserves_reviewer_edit_after_regeneration(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    update_review_item(project_dir, "cover-title", status="edited")
    set_queue_item(project_dir, "cover-title", edited_content="Reviewer edited front matter.")
    populate_for_review(project_dir)
    set_review_states(
        project_dir,
        overrides={"cover-title": {"status": "edited", "edited_content": "Reviewer edited front matter."}},
    )

    manifest = export_report(project_dir)
    markdown = Path(manifest["markdown_path"]).read_text(encoding="utf-8")

    assert "Reviewer edited front matter." in markdown


def test_cli_export_report_text_and_json(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    assert main(["export-report", str(project_dir), "--include-draft", "--format", "both"]) == 0
    captured = capsys.readouterr()
    assert "Generated preview export" in captured.out
    assert "DOCX:" in captured.out

    assert main(["export-report", str(project_dir), "--include-draft", "--json"]) == 0
    captured = capsys.readouterr()
    assert json.loads(captured.out)["project_id"] == "test_project"


def test_cli_export_report_gate_failure_text_and_json(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)

    assert main(["export-report", str(project_dir)]) == 1
    captured = capsys.readouterr()
    assert "review-complete gate" in captured.err
    assert "wetlands-and-waterbodies" in captured.err or "cover-title" in captured.err

    assert main(["export-report", str(project_dir), "--json"]) == 1
    captured = capsys.readouterr()
    payload = json.loads(captured.out)
    assert payload["review_gate_status"] == "blocked"
    assert payload["unreviewed_item_count"] > 0


def test_cli_export_report_qa_failure_text_and_json(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    project_dir = write_project(tmp_path)
    populate_for_review(project_dir)
    set_review_states(project_dir)
    set_queue_item(project_dir, "study-area", source_refs=["not_a_source_id"])

    assert main(["export-report", str(project_dir)]) == 1
    captured = capsys.readouterr()
    assert "export QA" in captured.err
    assert "unknown_source_ref" in captured.err

    assert main(["export-report", str(project_dir), "--json"]) == 1
    captured = capsys.readouterr()
    payload = json.loads(captured.out)
    assert payload["export_qa_status"] == "failed"
    assert payload["export_qa_blocking_error_count"] > 0
    assert any(issue["code"] == "unknown_source_ref" for issue in payload["export_qa_issues"])


def test_build_demo_deliverable_writes_package_manifest_without_accepting_items(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)

    manifest = build_demo_deliverable(project_dir, output_format="both")
    queue = load_review_queue(project_dir)

    assert manifest["package_status"] == "internal_preview"
    assert Path(manifest["output_path"]).exists()
    assert Path(manifest["markdown_path"]).exists()
    assert Path(manifest["docx_path"]).exists()
    assert "data_lineage" in manifest
    assert manifest["preview_mode"] is True
    assert manifest["review_gate_status"] == "preview_bypassed"
    assert manifest["unreviewed_item_count"] > 0
    assert manifest["final_verification"]["review_gate_status"] == "preview_bypassed"
    assert manifest["final_verification"]["docx_readable"] is True
    assert not any(item["status"] in {"accepted", "edited"} for item in queue["items"])


def test_build_mvp_deliverable_fails_without_real_source_data(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    project_dir = write_project(tmp_path)

    def populate_without_source_preparation(
        project_dir: Path,
        *,
        prepare_sources: bool,
        include_optional_sources: bool,
        materialize_local_sources: bool = False,
        gpt_drafting: bool | None = None,
        gpt_model: str | None = None,
    ) -> dict[str, Any]:
        return populate_for_review(project_dir, prepare_sources=False, include_optional_sources=False)

    monkeypatch.setattr(deliverable_module, "populate_for_review", populate_without_source_preparation)

    with pytest.raises(MvpDeliverableError, match="MVP deliverables require"):
        build_mvp_deliverable(project_dir)


def test_build_mvp_deliverable_fails_with_test_fixture_downloads(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    project_dir = write_project(tmp_path)
    monkeypatch.setattr(source_acquisition, "_fetch_json", fake_empty_public_fetch)

    with pytest.raises(MvpDeliverableError, match="test fixture"):
        build_mvp_deliverable(project_dir, fail_on_no_downloaded_sources=False)


def test_build_mvp_deliverable_succeeds_with_real_provided_source_layers(tmp_path: Path) -> None:
    project_dir = write_project(tmp_path)
    add_supported_real_source_inputs(project_dir)

    manifest = build_mvp_deliverable(project_dir, output_format="docx")
    text = docx_text(manifest["docx_path"])

    assert manifest["package_status"] == "internal_preview_real_data_mvp"
    assert manifest["preview_mode"] is True
    assert manifest["review_gate_status"] == "preview_bypassed"
    assert manifest["unreviewed_item_count"] > 0
    assert manifest["data_lineage"]["counts"]["provided_in_input"] >= 4
    assert manifest["data_lineage"]["counts"]["test_or_mock"] == 0
    assert manifest["mvp_quality"]["real_source_count"] >= 4
    assert manifest["mvp_quality"]["source_backed_constraint_count"] > 0
    assert manifest["final_verification"]["review_gate_status"] == "preview_bypassed"
    assert manifest["final_verification"]["docx_readable"] is True
    assert "Real Data Used" in text
    assert "provided_in_input" in json.dumps(manifest["data_lineage"])


def test_cli_build_mvp_deliverable_json_with_real_provided_sources(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    project_dir = write_project(tmp_path)
    add_supported_real_source_inputs(project_dir)

    assert main(["build-mvp-deliverable", str(project_dir), "--format", "docx", "--json"]) == 0
    captured = capsys.readouterr()
    payload = json.loads(captured.out)

    assert payload["project_id"] == "test_project"
    assert payload["data_lineage"]["real_source_count"] >= 4
    assert Path(payload["docx_path"]).exists()


def test_cli_build_demo_deliverable_json(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    project_dir = write_project(tmp_path)

    assert main(["build-demo-deliverable", str(project_dir), "--format", "docx", "--json"]) == 0
    captured = capsys.readouterr()
    payload = json.loads(captured.out)

    assert payload["project_id"] == "test_project"
    assert payload["markdown_path"] is None
    assert Path(payload["docx_path"]).exists()
