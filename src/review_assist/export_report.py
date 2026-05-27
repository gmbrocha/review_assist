"""Editable report export package generation from reviewed queue items."""

from __future__ import annotations

import json
import re
import shutil
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .data_lineage import build_data_lineage
from .deliverable_items import DeliverableItemsError, RENDER_POLICY_FIELDS, load_deliverable_items
from .deliverable_matrix import DeliverableMatrixError, REQUIRED_STUB_TEXT, load_deliverable_matrix
from .figure_style_model import FIGURE_VERSIONS_PATH, approved_figure_version
from .maps import MAP_MANIFEST_PATH, MapGenerationError, load_map_manifest
from .project_area import ProjectAreaError, load_project_area
from .projects import ProjectManifestError, load_project_manifest
from .report_section_policy import ReportSectionPolicyError, load_report_section_policy
from .review_queue import ReviewQueueError, generate_review_queue, load_review_queue
from .source_catalog import SourceCatalogError, load_source_catalog
from .source_status import SOURCE_STATUS_PATH, SourceStatusError, resolve_source_status_set
from .tables import TABLES_PATH, TableGenerationError, load_comparison_tables


EXPORT_DIR = Path("exports")
EXPORT_MANIFEST_PATH = EXPORT_DIR / "export_manifest.json"
EXPORT_MARKDOWN_PATH = EXPORT_DIR / "environmental_constraints_report.md"
EXPORT_DOCX_PATH = EXPORT_DIR / "environmental_constraints_report.docx"
EXPORT_FIGURE_ASSETS_DIR = EXPORT_DIR / "assets" / "figures"
SUPPORTED_OUTPUT_FORMATS = {"markdown", "docx", "both"}
FIGURE_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}
DOCX_TABLE_ROW_LIMIT = 50
DELIVERABLE_TABLE_BODY_PREVIEW_LIMIT = 5
EXPORT_BODY_CONTENT_WARNING_CHAR_LIMIT = 4000

EXPORT_GROUP_ORDER = [
    "front_matter",
    "executive_summary",
    "introduction",
    "methodology",
    "constraints_inventory",
    "natural_ecological",
    "cultural_historic",
    "community_resources",
    "utility_infrastructure",
    "contamination",
    "socioeconomic",
    "conclusion",
    "tables",
    "figures",
    "attachments",
    "resource_sections",
]
EXPORT_GROUP_TITLES = {
    "front_matter": "Front Matter",
    "executive_summary": "Executive Summary",
    "introduction": "Introduction",
    "methodology": "Methodology",
    "constraints_inventory": "Environmental Constraints Inventory",
    "natural_ecological": "Natural and Ecological Resources",
    "cultural_historic": "Cultural and Historic Resources",
    "community_resources": "Community Resources",
    "utility_infrastructure": "Utility and Infrastructure Considerations",
    "contamination": "Contamination Risks",
    "socioeconomic": "Socioeconomic and Business Considerations",
    "resource_sections": "Resource Sections",
    "conclusion": "Conclusion and Next Steps",
    "tables": "Tables",
    "figures": "Figures",
    "attachments": "Attachments",
}
UNRESOLVED_REQUIRED_SOURCE_STATUSES = {"missing", "downloadable", "needs_review", "gated", "stubbed", "failed"}
TERMINAL_REVIEW_STATUSES = {"accepted", "edited", "replaced", "declined"}
GATE_INCLUDED_NONTERMINAL_STATUSES = {"unable_to_verify"}
GATE_PREVIEW_LIMIT = 10
RAW_LEGACY_EXPORT_TYPES = {
    "draft_finding",
    "spatial_relationship",
    "comparison_table",
    "source_inventory_note",
    "source_status_note",
    "no_mapped_relationships",
    "validation_issue",
    "missing_data_placeholder",
}


class ExportReportError(RuntimeError):
    """Raised when editable report export generation cannot complete."""


class ExportGateError(ExportReportError):
    """Raised when default export is blocked by incomplete deliverable review."""

    def __init__(self, details: dict[str, Any]):
        self.details = details
        message = _review_gate_error_message(details)
        super().__init__(message)


class ExportQAError(ExportReportError):
    """Raised when reviewed export is blocked by policy-aware export QA."""

    def __init__(self, details: dict[str, Any]):
        self.details = details
        super().__init__(_export_qa_error_message(details))


def export_report(project_dir: Path, *, include_draft: bool = False, output_format: str = "markdown") -> dict[str, Any]:
    project_dir = project_dir.resolve()
    formats = _output_formats(output_format)
    try:
        queue = _load_or_generate_queue(project_dir)
        source_status = _load_or_generate_source_status(project_dir)
        comparison_tables = _load_optional_comparison_tables(project_dir)
        map_manifest = _load_optional_map_manifest(project_dir)
        matrix = load_deliverable_matrix()
        section_policy = load_report_section_policy()
        source_catalog = load_source_catalog()
    except (ReviewQueueError, SourceStatusError, DeliverableMatrixError, ReportSectionPolicyError, SourceCatalogError) as exc:
        raise ExportReportError(str(exc)) from exc

    output_dir = project_dir / EXPORT_DIR
    markdown_path = project_dir / EXPORT_MARKDOWN_PATH
    docx_path = project_dir / EXPORT_DOCX_PATH
    manifest_path = project_dir / EXPORT_MANIFEST_PATH
    now = _utc_now()

    items = _dict_list(queue.get("items", []))
    review_gate = _review_gate_summary(project_dir, queue, items, include_draft=include_draft)
    if not include_draft and review_gate["review_gate_status"] == "blocked":
        raise ExportGateError(review_gate)

    included, skipped = _partition_export_items(items, include_draft=include_draft)
    included = [*included, *_structural_heading_export_items(matrix, section_policy)]
    included = sorted(included, key=_export_sort_key)
    figure_version_issues = _resolve_export_figure_versions(project_dir, included)
    data_lineage = build_data_lineage(project_dir, included_items=included)
    unresolved_required_sources = _unresolved_required_sources(source_status)
    validation_issues = _export_validation_issues(
        included,
        unresolved_required_sources,
        include_draft=include_draft,
    )
    validation_issues.extend(_dict_list(review_gate.get("validation_issues", [])))
    validation_issues.extend(_dict_list(data_lineage.get("validation_issues", [])))
    if comparison_tables:
        validation_issues.extend(_dict_list(comparison_tables.get("validation_issues", [])))
    if map_manifest:
        validation_issues.extend(_dict_list(map_manifest.get("validation_issues", [])))
    validation_issues.extend(figure_version_issues)

    mvp_quality = _mvp_quality_summary(
        included=included,
        data_lineage=data_lineage,
        unresolved_required_sources=unresolved_required_sources,
        validation_issues=validation_issues,
        comparison_tables=comparison_tables,
        map_manifest=map_manifest,
        figure_assets=[],
    )
    validation_issues.extend(_mvp_quality_validation_issues(mvp_quality, include_draft=include_draft))
    mvp_quality = _mvp_quality_summary(
        included=included,
        data_lineage=data_lineage,
        unresolved_required_sources=unresolved_required_sources,
        validation_issues=validation_issues,
        comparison_tables=comparison_tables,
        map_manifest=map_manifest,
        figure_assets=[],
    )
    compactness_budget = _compactness_budget(
        included=included,
        review_gate=review_gate,
        mvp_quality=mvp_quality,
    )
    export_qa = _export_qa_summary(
        project_dir=project_dir,
        include_draft=include_draft,
        review_gate=review_gate,
        included=included,
        validation_issues=validation_issues,
        compactness_budget=compactness_budget,
        matrix=matrix,
        section_policy=section_policy,
        source_catalog=source_catalog,
        source_status=source_status,
        comparison_tables=comparison_tables,
        map_manifest=map_manifest,
    )
    if not include_draft and export_qa["export_qa_status"] == "failed":
        raise ExportQAError(export_qa)

    output_dir.mkdir(parents=True, exist_ok=True)
    figure_assets, figure_asset_issues = _prepare_export_figure_assets(project_dir, included, map_manifest)
    validation_issues.extend(figure_asset_issues)
    mvp_quality = _mvp_quality_summary(
        included=included,
        data_lineage=data_lineage,
        unresolved_required_sources=unresolved_required_sources,
        validation_issues=validation_issues,
        comparison_tables=comparison_tables,
        map_manifest=map_manifest,
        figure_assets=figure_assets,
    )
    compactness_budget = _compactness_budget(
        included=included,
        review_gate=review_gate,
        mvp_quality=mvp_quality,
    )

    if "markdown" in formats:
        markdown = _markdown_report(
            queue=queue,
            included=included,
            validation_issues=validation_issues,
            include_draft=include_draft,
            data_lineage=data_lineage,
            comparison_tables=comparison_tables,
            map_manifest=map_manifest,
            figure_assets=figure_assets,
        )
        markdown_path.write_text(markdown, encoding="utf-8")

    if "docx" in formats:
        _write_docx_report(
            project_dir=project_dir,
            docx_path=docx_path,
            queue=queue,
            included=included,
            validation_issues=validation_issues,
            include_draft=include_draft,
            comparison_tables=comparison_tables,
            map_manifest=map_manifest,
            data_lineage=data_lineage,
            figure_assets=figure_assets,
            output_paths={
                "markdown_report": str(markdown_path) if "markdown" in formats else None,
                "docx_report": str(docx_path),
                "export_manifest": str(manifest_path),
            },
        )

    final_verification = _final_verification_summary(
        included=included,
        review_gate=review_gate,
        compactness_budget=compactness_budget,
        output_formats=formats,
        markdown_path=markdown_path if "markdown" in formats else None,
        docx_path=docx_path if "docx" in formats else None,
    )
    validation_issues = _dedupe_issues([*validation_issues, *_dict_list(final_verification.get("issues", []))])
    export_qa = _export_qa_summary(
        project_dir=project_dir,
        include_draft=include_draft,
        review_gate=review_gate,
        included=included,
        validation_issues=validation_issues,
        compactness_budget=compactness_budget,
        matrix=matrix,
        section_policy=section_policy,
        source_catalog=source_catalog,
        source_status=source_status,
        comparison_tables=comparison_tables,
        map_manifest=map_manifest,
        final_verification=final_verification,
    )

    manifest = {
        "project_id": queue.get("project_id"),
        "project_name": queue.get("project_name"),
        "project_dir": str(project_dir),
        "created_at": now,
        "include_draft": include_draft,
        "preview_mode": include_draft,
        "output_format": output_format,
        "output_formats": formats,
        "package_status": "internal_preview" if include_draft else "reviewed_content",
        "review_gate_status": review_gate["review_gate_status"],
        "review_gate": review_gate,
        "export_qa_status": export_qa["export_qa_status"],
        "export_qa_issue_count": export_qa["export_qa_issue_count"],
        "export_qa_blocking_error_count": export_qa["export_qa_blocking_error_count"],
        "export_qa_issues": export_qa["export_qa_issues"],
        "export_qa": export_qa,
        "review_queue_path": queue.get("output_path"),
        "source_status_path": source_status.get("output_path"),
        "comparison_tables_path": comparison_tables.get("output_path") if comparison_tables else None,
        "markdown_path": str(markdown_path) if "markdown" in formats else None,
        "docx_path": str(docx_path) if "docx" in formats else None,
        "included_count": len(included),
        "skipped_count": len(skipped),
        "included_item_count": len(included),
        "skipped_item_count": len(skipped),
        "review_item_count": review_gate["review_item_count"],
        "terminal_review_item_count": review_gate["terminal_review_item_count"],
        "unreviewed_item_count": review_gate["unreviewed_item_count"],
        "declined_item_count": review_gate["declined_item_count"],
        "deliverable_matrix_version": review_gate.get("deliverable_matrix_version", ""),
        "expected_deliverable_item_count": review_gate.get("expected_deliverable_item_count"),
        "actual_deliverable_item_count": review_gate.get("actual_deliverable_item_count"),
        "stub_item_count": review_gate["stub_item_count"],
        "unreviewed_items_preview": review_gate["unreviewed_items_preview"],
        "status_counts": dict(Counter(str(item.get("status", "")) for item in items)),
        "included_status_counts": dict(Counter(item["status"] for item in included)),
        "skipped_status_counts": dict(Counter(item["status"] for item in skipped)),
        "included_type_counts": dict(Counter(item["type"] for item in included)),
        "skipped_type_counts": dict(Counter(item["type"] for item in skipped)),
        "included_table_ids": sorted({str(item.get("table_id")) for item in included if item.get("table_id")}),
        "included_figure_ids": sorted(_included_figure_ids(included)),
        "included_attachment_ids": sorted(_included_attachment_ids(included)),
        "included_map_paths": sorted({_export_map_path(item) for item in included if _export_map_path(item)}),
        "export_figure_assets": figure_assets,
        "included_source_refs": sorted({ref for item in included for ref in _string_list(item.get("source_refs", []))}),
        "unresolved_required_sources": unresolved_required_sources,
        "data_lineage": data_lineage,
        "evidence_package_path": _evidence_package_path(queue),
        "gpt_drafting": _gpt_drafting_summary(items),
        "mvp_quality": mvp_quality,
        "compactness_budget": compactness_budget,
        "final_verification": final_verification,
        "package_contents": _package_contents(queue, output_paths={
            "markdown_report": str(markdown_path) if "markdown" in formats else None,
            "docx_report": str(docx_path) if "docx" in formats else None,
            "export_manifest": str(manifest_path),
        }, figure_assets=figure_assets),
        "included_items": included,
        "skipped_items": skipped,
        "validation_issues": validation_issues,
        "output_path": str(manifest_path),
    }
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return manifest


def _load_or_generate_queue(project_dir: Path) -> dict[str, Any]:
    try:
        return load_review_queue(project_dir)
    except ReviewQueueError:
        return generate_review_queue(project_dir)


def _load_or_generate_source_status(project_dir: Path) -> dict[str, Any]:
    source_status_path = project_dir / SOURCE_STATUS_PATH
    if source_status_path.exists():
        try:
            data = json.loads(source_status_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError as exc:
            raise SourceStatusError(f"Invalid source status JSON: {source_status_path}: {exc}") from exc
        if isinstance(data, dict):
            return data
        raise SourceStatusError(f"Source status artifact must be a JSON object: {source_status_path}")
    return resolve_source_status_set(project_dir)


def _load_optional_comparison_tables(project_dir: Path) -> dict[str, Any] | None:
    tables_path = project_dir / TABLES_PATH
    if not tables_path.exists():
        return None
    try:
        return load_comparison_tables(project_dir)
    except TableGenerationError as exc:
        return {
            "tables": [],
            "output_path": str(tables_path),
            "validation_issues": [
                _issue(
                    "warning",
                    "comparison_tables_unavailable",
                    f"Comparison table artifact could not be loaded for export table rendering: {exc}",
                )
            ],
        }


def _load_optional_map_manifest(project_dir: Path) -> dict[str, Any] | None:
    manifest_path = project_dir / MAP_MANIFEST_PATH
    if not manifest_path.exists():
        return None
    try:
        return load_map_manifest(project_dir)
    except MapGenerationError as exc:
        return {
            "figures": [],
            "figure_count": 0,
            "output_path": str(manifest_path),
            "validation_issues": [
                _issue(
                    "warning",
                    "map_manifest_unavailable",
                    f"Map manifest artifact could not be loaded for export figure rendering: {exc}",
                )
            ],
        }


def _prepare_export_figure_assets(
    project_dir: Path,
    included: list[dict[str, Any]],
    map_manifest: dict[str, Any] | None,
) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    figure_lookup = _figures_by_id(map_manifest, included)
    included_figure_ids = sorted(_included_figure_ids(included))
    if not included_figure_ids:
        return [], []
    output_dir = project_dir / EXPORT_FIGURE_ASSETS_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    assets: list[dict[str, Any]] = []
    issues: list[dict[str, str]] = []
    for figure_id in included_figure_ids:
        figure = figure_lookup.get(figure_id, {})
        source_value = figure.get("image_path")
        if not source_value:
            issues.append(_issue("warning", "missing_export_figure_source", f"Included figure '{figure_id}' has no image path to copy."))
            continue
        source_path = _resolve_project_path(project_dir, source_value)
        if not source_path.exists():
            issues.append(_issue("warning", "missing_export_figure_asset", f"Included figure '{figure_id}' image file was not found for export asset copy: {source_path}."))
            continue
        destination = output_dir / f"{_slug(figure_id)}{source_path.suffix or '.png'}"
        try:
            shutil.copy2(source_path, destination)
        except OSError as exc:
            issues.append(_issue("warning", "export_figure_asset_copy_failed", f"Included figure '{figure_id}' could not be copied into the export package: {exc}."))
            continue
        relative_asset = destination.relative_to(project_dir / EXPORT_DIR).as_posix()
        asset = {
            "figure_id": figure_id,
            "title": figure.get("title") or figure_id,
            "source_image_path": str(source_path),
            "export_image_path": str(destination),
            "export_asset_path": relative_asset,
            **_selected_figure_version_fields(figure),
        }
        assets.append(asset)
        _annotate_figure_asset(included, figure_lookup, figure_id, asset)
    return assets, issues


def _resolve_export_figure_versions(project_dir: Path, included: list[dict[str, Any]]) -> list[dict[str, str]]:
    """Apply approved/regenerated figure-version image selection to export items."""

    versions_artifact = _load_optional_figure_versions(project_dir)
    issues: list[dict[str, str]] = []
    for item in included:
        if item.get("type") not in {"map_figure", "figure"}:
            continue
        figure_id = _item_figure_id(item)
        if not figure_id or bool(item.get("is_stub", False)):
            continue
        if str(item.get("image_source") or "") == "replacement_figure":
            item.update(_implicit_figure_version_metadata("replacement_figure", str(item.get("image_path") or "")))
            continue
        selected = _select_export_figure_version(versions_artifact, figure_id)
        if selected is None:
            item.update(_implicit_figure_version_metadata("current_generated_image", str(item.get("image_path") or "")))
            continue
        path = str(selected.get("output_artifact_path") or "").strip()
        source = _version_export_source(selected)
        item.update(_version_export_metadata(selected, source))
        if path:
            item["image_path"] = path
        if source in {"approved_version", "regenerated_version"}:
            if not path:
                issues.append(
                    _qa_issue(
                        "error",
                        "selected_figure_version_path_missing",
                        f"Selected {source.replace('_', ' ')} for figure '{figure_id}' has no output artifact path.",
                        item_id=str(item.get("id") or ""),
                        target_id=figure_id,
                    )
                )
            elif not _resolve_project_path(project_dir, path).exists():
                issues.append(
                    _qa_issue(
                        "error",
                        "selected_figure_version_asset_missing",
                        f"Selected {source.replace('_', ' ')} for figure '{figure_id}' is missing: {path}.",
                        item_id=str(item.get("id") or ""),
                        target_id=figure_id,
                    )
                )
    return issues


def _load_optional_figure_versions(project_dir: Path) -> dict[str, Any]:
    path = project_dir / FIGURE_VERSIONS_PATH
    if not path.exists():
        return {"version_count": 0, "versions": []}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {"version_count": 0, "versions": []}
    return data if isinstance(data, dict) else {"version_count": 0, "versions": []}


def _select_export_figure_version(versions_artifact: dict[str, Any], figure_id: str) -> dict[str, Any] | None:
    approved = approved_figure_version(versions_artifact, figure_id)
    if approved is not None:
        return approved
    versions = [
        version
        for version in _dict_list(versions_artifact.get("versions", []))
        if str(version.get("figure_id") or "") == figure_id
    ]
    regenerated = [
        version
        for version in versions
        if str(version.get("source_image_status") or "") == "regenerated_review_only"
        or str(version.get("approval_state") or "") == "regenerated"
    ]
    if regenerated:
        return sorted(regenerated, key=_figure_version_sort_key)[-1]
    autogenerated = [version for version in versions if str(version.get("approval_state") or "") == "autogenerated"]
    if autogenerated:
        return sorted(autogenerated, key=_figure_version_sort_key)[-1]
    return None


def _version_export_metadata(version: dict[str, Any], source: str) -> dict[str, Any]:
    return {
        "selected_figure_version_id": str(version.get("version_id") or ""),
        "selected_figure_version_number": version.get("version_number"),
        "selected_figure_version_state": str(version.get("approval_state") or ""),
        "selected_figure_version_source": source,
        "selected_figure_version_path": str(version.get("output_artifact_path") or ""),
    }


def _implicit_figure_version_metadata(source: str, path: str) -> dict[str, Any]:
    return {
        "selected_figure_version_id": "",
        "selected_figure_version_number": None,
        "selected_figure_version_state": "",
        "selected_figure_version_source": source,
        "selected_figure_version_path": path,
    }


def _selected_figure_version_fields(figure: dict[str, Any]) -> dict[str, Any]:
    keys = (
        "selected_figure_version_id",
        "selected_figure_version_number",
        "selected_figure_version_state",
        "selected_figure_version_source",
        "selected_figure_version_path",
    )
    return {key: figure.get(key) for key in keys if key in figure}


def _version_export_source(version: dict[str, Any]) -> str:
    state = str(version.get("approval_state") or "")
    if state == "approved":
        return "approved_version"
    if state == "regenerated" or str(version.get("source_image_status") or "") == "regenerated_review_only":
        return "regenerated_version"
    if state == "autogenerated":
        return "autogenerated_version"
    return "figure_version"


def _figure_version_sort_key(version: dict[str, Any]) -> tuple[int, str, str]:
    try:
        number = int(version.get("version_number") or 0)
    except (TypeError, ValueError):
        number = 0
    timestamp = str(version.get("updated_at") or version.get("completed_at") or version.get("created_at") or "")
    return number, timestamp, str(version.get("version_id") or "")


def _annotate_figure_asset(
    included: list[dict[str, Any]],
    figure_lookup: dict[str, dict[str, Any]],
    figure_id: str,
    asset: dict[str, Any],
) -> None:
    if figure_id in figure_lookup:
        figure_lookup[figure_id].update(asset)
    for item in included:
        if _item_figure_id(item) == figure_id:
            item.update(asset)


def _resolve_project_path(project_dir: Path, value: Any) -> Path:
    path = Path(str(value))
    return path if path.is_absolute() else project_dir / path


def _output_formats(output_format: str) -> list[str]:
    if output_format not in SUPPORTED_OUTPUT_FORMATS:
        allowed = ", ".join(sorted(SUPPORTED_OUTPUT_FORMATS))
        raise ExportReportError(f"Unsupported export format '{output_format}'. Expected one of: {allowed}.")
    if output_format == "both":
        return ["markdown", "docx"]
    return [output_format]


def _review_gate_summary(
    project_dir: Path,
    queue: dict[str, Any],
    items: list[dict[str, Any]],
    *,
    include_draft: bool,
) -> dict[str, Any]:
    queue_mode = str(queue.get("queue_mode") or "")
    metadata = _deliverable_items_metadata(project_dir, queue)
    validation_issues: list[dict[str, Any]] = []
    unreviewed: list[dict[str, Any]] = []
    terminal_count = 0
    declined_count = 0
    export_includable_nonterminal_count = 0

    if queue_mode != "deliverable_items":
        issue = _issue(
            "error",
            "standard_review_queue_required",
            "Default export requires the standard matrix-bounded deliverable review queue. Regenerate the queue without legacy/audit mode.",
        )
        validation_issues.append(issue)
        unreviewed.append(
            {
                "id": "",
                "title": "Standard deliverable review queue required",
                "status": queue_mode or "unknown",
                "reason": "standard_review_queue_required",
            }
        )

    for item in items:
        status = _normalized_status(item.get("status"))
        item_id = str(item.get("id") or item.get("deliverable_item_id") or item.get("target_id") or "")
        if status in TERMINAL_REVIEW_STATUSES:
            if status == "declined":
                declined_count += 1
            if status == "replaced" and not _replacement_content(item):
                unreviewed.append(_gate_item(item, "replacement_content_missing"))
                validation_issues.append(_gate_issue("replacement_content_missing", item_id, "Review item is marked replaced but has no replacement content."))
            else:
                terminal_count += 1
                if status == "edited" and not _edited_content(item):
                    validation_issues.append(_gate_issue("edited_content_missing", item_id, "Review item is marked edited but has no edited content; export will fall back to generated content."))
            continue
        if status in GATE_INCLUDED_NONTERMINAL_STATUSES and _is_explicitly_export_includable(item):
            export_includable_nonterminal_count += 1
            continue
        unreviewed.append(_gate_item(item, _gate_block_reason(item)))

    if not items:
        validation_issues.append(_issue("error", "no_review_queue_items", "Default export requires at least one standard deliverable review item."))
        unreviewed.append({"id": "", "title": "No review queue items", "status": "missing", "reason": "no_review_queue_items"})

    status = "preview_bypassed" if include_draft else ("blocked" if unreviewed else "passed")
    return {
        "review_gate_status": status,
        "preview_mode": include_draft,
        "queue_mode": queue_mode,
        "review_item_count": len(items),
        "terminal_review_item_count": terminal_count,
        "export_includable_nonterminal_count": export_includable_nonterminal_count,
        "unreviewed_item_count": len(unreviewed),
        "declined_item_count": declined_count,
        "stub_item_count": _stub_item_count(items),
        "deliverable_matrix_version": metadata.get("matrix_version", ""),
        "expected_deliverable_item_count": metadata.get("expected_item_count", len(items)),
        "actual_deliverable_item_count": metadata.get("item_count", len(items)),
        "unreviewed_items_preview": unreviewed[:GATE_PREVIEW_LIMIT],
        "validation_issues": _dedupe_issues(validation_issues),
        "message": (
            "Internal preview bypassed the review-complete export gate."
            if include_draft
            else (
                "Review-complete export gate passed."
                if not unreviewed
                else "Default export is blocked until all standard deliverable review items are accepted, edited, replaced, declined, or explicitly export-eligible unable-to-verify items."
            )
        ),
    }


def _deliverable_items_metadata(project_dir: Path, queue: dict[str, Any]) -> dict[str, Any]:
    upstream = queue.get("upstream_artifacts", {}) if isinstance(queue.get("upstream_artifacts"), dict) else {}
    path_value = upstream.get("deliverable_items_path")
    if path_value:
        path = Path(str(path_value))
        if not path.is_absolute():
            path = project_dir / path
        if path.exists():
            try:
                artifact = load_deliverable_items(project_dir)
                return {
                    "matrix_version": artifact.get("matrix_version", ""),
                    "expected_item_count": artifact.get("expected_item_count"),
                    "item_count": artifact.get("item_count"),
                }
            except DeliverableItemsError:
                pass
    return _metadata_from_queue_items(queue)


def _metadata_from_queue_items(queue: dict[str, Any]) -> dict[str, Any]:
    items = _dict_list(queue.get("items", []))
    matrix_version = ""
    for item in items:
        provenance = item.get("provenance", {}) if isinstance(item.get("provenance"), dict) else {}
        deliverable_provenance = provenance.get("deliverable_item_provenance", {}) if isinstance(provenance.get("deliverable_item_provenance"), dict) else {}
        matrix_version = str(deliverable_provenance.get("matrix_version") or "")
        if matrix_version:
            break
    return {
        "matrix_version": matrix_version,
        "expected_item_count": len(items),
        "item_count": len(items),
    }


def _gate_item(item: dict[str, Any], reason: str) -> dict[str, str]:
    return {
        "id": str(item.get("id") or item.get("deliverable_item_id") or item.get("target_id") or ""),
        "title": str(item.get("title") or ""),
        "status": _normalized_status(item.get("status")),
        "reason": reason,
    }


def _gate_block_reason(item: dict[str, Any]) -> str:
    status = _normalized_status(item.get("status"))
    if status == "unable_to_verify" and item.get("export_eligible") is not True:
        return "unable_to_verify_not_export_eligible"
    if status == "unable_to_verify" and not _best_export_content(item):
        return "unable_to_verify_missing_content"
    return f"status_{status or 'unknown'}"


def _gate_issue(code: str, item_id: str, message: str) -> dict[str, str]:
    return {
        "severity": "warning",
        "code": code,
        "message": message,
        "location": "review_queue/review_queue.json",
        "item_id": item_id,
    }


def _review_gate_error_message(details: dict[str, Any]) -> str:
    preview = details.get("unreviewed_items_preview", [])
    examples = ", ".join(
        f"{item.get('id') or 'unknown'} [{item.get('status')}]"
        for item in _dict_list(preview)
        if item.get("id") or item.get("status")
    )
    suffix = f" First unreviewed items: {examples}." if examples else ""
    return (
        "Default export is blocked by the review-complete gate: "
        f"{details.get('unreviewed_item_count', 0)} of {details.get('review_item_count', 0)} deliverable review item(s) are not export-ready."
        f"{suffix} Use --include-draft only for an internal/pre-review preview."
    )


def _include_item(item: dict[str, Any], *, include_draft: bool) -> bool:
    status = _normalized_status(item.get("status"))
    if include_draft:
        return status != "declined"
    if _report_body_ineligible(item):
        return (
            status in {"edited", "replaced", "unable_to_verify"}
            and bool(item.get("export_eligible", False))
            and bool(_reviewer_supplied_export_content(item))
        )
    if status in {"accepted", "edited"}:
        return bool(item.get("export_eligible", False))
    if status == "replaced":
        return bool(item.get("export_eligible", False)) and bool(_replacement_content(item))
    if status == "unable_to_verify":
        return _is_explicitly_export_includable(item)
    return False


def _partition_export_items(
    items: list[dict[str, Any]],
    *,
    include_draft: bool,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    included: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    for item in items:
        if _include_item(item, include_draft=include_draft):
            included.append(_export_item(item))
        else:
            skipped.append(_skipped_item(item, include_draft=include_draft))
    return included, skipped


def _structural_heading_export_items(matrix: Any, section_policy: Any) -> list[dict[str, Any]]:
    policies = section_policy.by_section_id()
    items: list[dict[str, Any]] = []
    for target in getattr(matrix, "section_targets", []):
        policy = policies.get(str(target.target_id))
        if policy is None or policy.section_role != "structural_heading":
            continue
        render_policy = {
            "policy_inclusion_status": policy.inclusion_status,
            "policy_activation_condition": policy.activation_condition,
            "policy_review_requirement": policy.review_requirement,
            "policy_comparison_unit_expansion": policy.comparison_unit_expansion_policy,
            "render_decision": "include_body",
            "render_destination": "report_outline",
            "render_decision_reason": "Structural heading is emitted from matrix policy and does not require review.",
            "report_body_eligible": True,
        }
        items.append(
            {
                "id": str(target.target_id),
                "target_id": str(target.target_id),
                "deliverable_item_id": str(target.target_id),
                "type": "structural_heading",
                "title": str(target.title),
                "status": "not_review_required",
                "is_stub": False,
                "export_group": str(target.export_group),
                "export_section": str(target.target_id),
                "section_number": target.section_number,
                "section_order": _optional_number(target.section_order),
                "heading_level": _optional_int(target.heading_level),
                "content": "",
                "content_source": "structural_heading",
                "source_refs": [],
                "uncertainty_flags": [],
                "required_caveats": [],
                "allowed_source_refs": [],
                "allowed_source_categories": [],
                "allowed_table_refs": [],
                "allowed_figure_refs": [],
                "visual_slots": [],
                "table_slots": [],
                "related_figure_ids": [],
                "related_table_ids": [],
                "image_path": "",
                "figure_id": "",
                "figure_type": "",
                "caption": "",
                "caption_source": "",
                "image_source": "",
                "source_note": "",
                "method_note": "",
                "map_elements": [],
                "figure_group": "",
                "related_resource_categories": [],
                "table_id": "",
                "columns": [],
                "rows_preview": [],
                "required_columns": [],
                "row_count": None,
                "attachment_id": "",
                "attachment_refs": [],
                "artifact_path": "",
                "provenance": {
                    "artifact": "deliverable_matrix",
                    "matrix_version": getattr(matrix, "matrix_version", ""),
                    "section_role": "structural_heading",
                    "review_before_export": False,
                },
                "manual_material": {
                    "material_type": "none",
                    "material_status": "not_used",
                    "export_behavior": "do_not_export",
                    "reviewer_action": "No review action is required for structural headings.",
                    "source_refs": [],
                    "source_categories": [],
                    "internal_note_only": False,
                },
                **render_policy,
                "render_policy": render_policy,
            }
        )
    return items


def _export_item(item: dict[str, Any]) -> dict[str, Any]:
    status = _normalized_status(item.get("status"))
    assumptions = item.get("assumptions", {}) if isinstance(item.get("assumptions"), dict) else {}
    matrix_target = assumptions.get("matrix_target", {}) if isinstance(assumptions.get("matrix_target"), dict) else {}
    provenance = item.get("provenance", {}) if isinstance(item.get("provenance"), dict) else {}
    generated = _generated_content(item)
    edited = _edited_content(item)
    replacement = _replacement_content(item)
    if _is_figure_item(item):
        caption, caption_source = _figure_caption(item, assumptions)
        image_path, image_source = _figure_image_path(item, assumptions)
        if image_path:
            content = ""
            content_source = image_source if image_source == "replacement_figure" else caption_source
        elif status == "replaced" and replacement and not _is_figure_image_path(replacement):
            content = replacement
            content_source = "replacement_content"
        elif status == "edited" and edited:
            content = edited
            content_source = "edited_caption"
        else:
            content = generated
            content_source = "generated_content"
    elif status == "replaced":
        content = replacement
        content_source = "replacement_content"
    elif status == "edited" and edited:
        content = edited
        content_source = "edited_content"
    elif status == "unable_to_verify":
        content = _best_export_content(item)
        content_source = "replacement_content" if replacement else ("edited_content" if edited else "generated_content")
    else:
        content = generated
        content_source = "generated_content"
    if not _is_figure_item(item):
        caption = item.get("caption") or assumptions.get("caption")
        caption_source = "generated_caption"
        image_path = item.get("image_path") or assumptions.get("image_path")
        image_source = "generated_figure"
    render_policy = _render_policy_fields(item)
    manual_material = _manual_material_fields(item)
    return {
        "id": str(item.get("id", "")),
        "target_id": str(item.get("target_id") or item.get("id", "")),
        "deliverable_item_id": str(item.get("deliverable_item_id") or item.get("id", "")),
        "type": str(item.get("type", "")),
        "title": str(item.get("title", "")),
        "status": status,
        "is_stub": bool(item.get("is_stub", False)),
        "export_group": str(item.get("export_group") or _default_export_group(item)),
        "export_section": str(item.get("export_section", "")),
        "section_number": item.get("section_number") or matrix_target.get("section_number"),
        "section_order": _optional_number(item.get("section_order") if item.get("section_order") is not None else assumptions.get("section_order")),
        "heading_level": _optional_int(item.get("heading_level") or assumptions.get("heading_level") or matrix_target.get("heading_level")),
        "content": content,
        "content_source": content_source,
        "source_refs": _string_list(item.get("source_refs", [])),
        "uncertainty_flags": _string_list(item.get("uncertainty_flags", [])),
        "required_caveats": _policy_list_from_item(item, assumptions, "required_caveats"),
        "allowed_source_refs": _policy_list_from_item(item, assumptions, "allowed_source_refs"),
        "allowed_source_categories": _policy_list_from_item(item, assumptions, "allowed_source_categories"),
        "allowed_table_refs": _policy_list_from_item(item, assumptions, "allowed_table_refs"),
        "allowed_figure_refs": _policy_list_from_item(item, assumptions, "allowed_figure_refs"),
        "visual_slots": _string_list(assumptions.get("visual_slots", [])),
        "table_slots": _string_list(assumptions.get("table_slots", [])),
        "related_figure_ids": _string_list(item.get("related_figure_ids", [])),
        "related_table_ids": _string_list(item.get("related_table_ids", [])),
        "image_path": image_path,
        "figure_id": item.get("figure_id"),
        "figure_type": item.get("figure_type"),
        "caption": caption,
        "caption_source": caption_source,
        "image_source": image_source,
        "source_note": item.get("source_note") or assumptions.get("source_note"),
        "method_note": item.get("method_note") or assumptions.get("method_note"),
        "map_elements": _string_list(item.get("map_elements", [])),
        "figure_group": item.get("figure_group"),
        "related_resource_categories": _string_list(item.get("related_resource_categories", [])),
        "table_id": item.get("table_id"),
        "columns": _string_list(item.get("columns", [])) or _string_list(assumptions.get("columns", [])),
        "rows_preview": _dict_list(item.get("rows_preview", [])) or _dict_list(assumptions.get("rows_preview", [])),
        "required_columns": _string_list(item.get("required_columns", [])) or _string_list(assumptions.get("required_columns", [])),
        "row_count": item.get("row_count") if item.get("row_count") is not None else assumptions.get("row_count"),
        "attachment_id": item.get("attachment_id"),
        "attachment_refs": _string_list(item.get("attachment_refs", [])) or _string_list(matrix_target.get("attachment_refs", [])),
        "artifact_path": provenance.get("artifact_path"),
        "provenance": provenance,
        "manual_material": manual_material,
        **render_policy,
        "render_policy": render_policy,
    }


def _is_figure_item(item: dict[str, Any]) -> bool:
    return str(item.get("type") or "") in {"map_figure", "figure"} or bool(str(item.get("figure_id") or "").strip())


def _figure_caption(item: dict[str, Any], assumptions: dict[str, Any]) -> tuple[str, str]:
    figure_review = item.get("figure_review", {}) if isinstance(item.get("figure_review"), dict) else {}
    reviewed_caption = str(figure_review.get("caption") or "").strip()
    if reviewed_caption:
        return reviewed_caption, str(figure_review.get("caption_source") or "reviewed_caption")
    edited = _edited_content(item)
    if edited:
        return edited, "edited_caption"
    return str(item.get("caption") or assumptions.get("caption") or "").strip(), "generated_caption"


def _figure_image_path(item: dict[str, Any], assumptions: dict[str, Any]) -> tuple[Any, str]:
    figure_review = item.get("figure_review", {}) if isinstance(item.get("figure_review"), dict) else {}
    reviewed_image = str(figure_review.get("image_path") or "").strip()
    if reviewed_image:
        return reviewed_image, str(figure_review.get("image_source") or "reviewed_figure")
    replacement = _replacement_content(item)
    if _is_figure_image_path(replacement):
        return replacement, "replacement_figure"
    return item.get("image_path") or assumptions.get("image_path"), "generated_figure"


def _is_figure_image_path(value: Any) -> bool:
    text = str(value or "").strip()
    return bool(text) and Path(text).suffix.lower() in FIGURE_IMAGE_EXTENSIONS


def _skipped_item(item: dict[str, Any], *, include_draft: bool) -> dict[str, Any]:
    render_policy = _render_policy_fields(item)
    manual_material = _manual_material_fields(item)
    return {
        "id": str(item.get("id", "")),
        "type": str(item.get("type", "")),
        "title": str(item.get("title", "")),
        "status": _normalized_status(item.get("status")),
        "export_group": str(item.get("export_group") or _default_export_group(item)),
        "reason": _skip_reason(item, include_draft=include_draft),
        "manual_material": manual_material,
        **render_policy,
        "render_policy": render_policy,
    }


def _skip_reason(item: dict[str, Any], *, include_draft: bool) -> str:
    status = _normalized_status(item.get("status"))
    if include_draft and status == "declined":
        return "declined"
    if status == "replaced" and not _replacement_content(item):
        return "replacement_content_missing"
    if status == "unable_to_verify" and not _best_export_content(item):
        return "unable_to_verify_missing_content"
    if not include_draft and _report_body_ineligible(item):
        if _reviewer_supplied_export_content(item) and not item.get("export_eligible", False):
            return "not_export_eligible"
        return f"policy_render_{str(item.get('render_decision') or 'body_ineligible')}"
    if status in {"accepted", "edited", "replaced", "unable_to_verify"} and not item.get("export_eligible", False):
        return "not_export_eligible"
    return f"status_{status or 'unknown'}"


def _report_body_ineligible(item: dict[str, Any]) -> bool:
    return str(item.get("type") or "") in {"front_matter", "report_section", "section_text"} and _coerce_bool(
        item.get("report_body_eligible", True)
    ) is False


def _reviewer_supplied_export_content(item: dict[str, Any]) -> str:
    return _replacement_content(item) or _edited_content(item)


def _render_policy_fields(item: dict[str, Any]) -> dict[str, Any]:
    defaults = {
        "policy_inclusion_status": "default",
        "policy_activation_condition": "always",
        "policy_review_requirement": "standard_review",
        "policy_comparison_unit_expansion": "none",
        "render_decision": "include_body",
        "render_destination": "report_body",
        "render_decision_reason": "No render gating applies.",
        "report_body_eligible": True,
    }
    result: dict[str, Any] = {}
    for key in RENDER_POLICY_FIELDS:
        value = item.get(key, defaults[key])
        result[key] = _coerce_bool(value) if key == "report_body_eligible" else str(value)
    return result


def _manual_material_fields(item: dict[str, Any]) -> dict[str, Any]:
    record = item.get("manual_material", {}) if isinstance(item.get("manual_material"), dict) else {}
    result = {
        "material_type": str(record.get("material_type") or "none"),
        "material_status": str(record.get("material_status") or "not_used"),
        "export_behavior": str(record.get("export_behavior") or "do_not_export"),
        "reviewer_action": str(record.get("reviewer_action") or ""),
        "source_refs": _string_list(item.get("source_refs", [])) or _string_list(record.get("source_refs", [])),
        "source_categories": _string_list(record.get("source_categories", [])),
        "internal_note_only": bool(record.get("internal_note_only", False)),
    }
    status = _normalized_status(item.get("status"))
    edited = _edited_content(item)
    replacement = _replacement_content(item)
    is_figure = _is_figure_item(item)
    if status == "declined":
        result["material_status"] = "not_used"
        result["export_behavior"] = "do_not_export"
    elif status == "unable_to_verify":
        result["material_status"] = "unable_to_verify"
    elif replacement:
        result["material_status"] = "reviewer_supplied"
        result["material_type"] = "replacement_figure" if is_figure and _is_figure_image_path(replacement) else "manual_text"
        result["export_behavior"] = "figure_review_when_reviewed" if is_figure else "body_replacement_when_reviewed"
    elif edited:
        result["material_status"] = "reviewer_supplied"
        result["material_type"] = "edited_caption" if is_figure else "manual_text"
        result["export_behavior"] = "figure_review_when_reviewed" if is_figure else "body_replacement_when_reviewed"
    return result


def _policy_list_from_item(item: dict[str, Any], assumptions: dict[str, Any], key: str) -> list[str]:
    values = _string_list(item.get(key, []))
    if values:
        return values
    values = _string_list(assumptions.get(key, []))
    if values:
        return values
    matrix_target = assumptions.get("matrix_target", {}) if isinstance(assumptions.get("matrix_target"), dict) else {}
    return _string_list(matrix_target.get(key, []))


def _coerce_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes"}
    return bool(value)


def _normalized_status(value: Any) -> str:
    status = str(value or "").strip()
    if status == "rejected":
        return "declined"
    return status


def _generated_content(item: dict[str, Any]) -> str:
    return str(item.get("generated_content") or "").strip()


def _edited_content(item: dict[str, Any]) -> str:
    return str(item.get("edited_content") or "").strip()


def _replacement_content(item: dict[str, Any]) -> str:
    return str(item.get("replacement_content") or "").strip()


def _best_export_content(item: dict[str, Any]) -> str:
    return _replacement_content(item) or _edited_content(item) or _generated_content(item)


def _is_explicitly_export_includable(item: dict[str, Any]) -> bool:
    return bool(item.get("export_eligible", False)) and bool(_best_export_content(item))


def _export_sort_key(item: dict[str, Any]) -> tuple[int, int, int, str]:
    group = str(item.get("export_group", ""))
    group_index = EXPORT_GROUP_ORDER.index(group) if group in EXPORT_GROUP_ORDER else len(EXPORT_GROUP_ORDER)
    item_type = str(item.get("type", ""))
    type_index = 0 if item_type in {"report_section", "section_text", "structural_heading"} else 1
    order = _sort_order_value(item.get("section_order"))
    return (group_index, type_index, order, str(item.get("title", "")))


def _sort_order_value(value: Any) -> int:
    if isinstance(value, bool):
        return 9999
    if isinstance(value, (int, float)):
        return int(value * 1000)
    try:
        return int(float(value) * 1000)
    except (TypeError, ValueError):
        return 9999


def _export_validation_issues(
    included: list[dict[str, Any]],
    unresolved_required_sources: list[dict[str, Any]],
    *,
    include_draft: bool,
) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    if not include_draft and not any(item["type"] in {"report_section", "section_text"} for item in included):
        issues.append(_issue("warning", "no_accepted_report_sections", "No accepted or edited report sections were available for export."))
    if not include_draft and not any(item["type"] in {"map_figure", "figure"} for item in included):
        issues.append(_issue("warning", "no_accepted_maps", "No accepted or edited map figures were available for export."))
    if unresolved_required_sources:
        categories = ", ".join(str(item.get("category")) for item in unresolved_required_sources)
        issues.append(
            _issue(
                "warning",
                "unresolved_required_source_gaps",
                f"Required source categories remain unresolved or require review: {categories}.",
            )
        )
    return issues


def _export_qa_summary(
    *,
    project_dir: Path,
    include_draft: bool,
    review_gate: dict[str, Any],
    included: list[dict[str, Any]],
    validation_issues: list[dict[str, Any]],
    compactness_budget: dict[str, Any],
    matrix: Any,
    section_policy: Any,
    source_catalog: Any,
    source_status: dict[str, Any],
    comparison_tables: dict[str, Any] | None,
    map_manifest: dict[str, Any] | None,
    final_verification: dict[str, Any] | None = None,
) -> dict[str, Any]:
    table_ids = {str(target.target_id) for target in getattr(matrix, "table_targets", [])}
    figure_ids = {str(target.target_id) for target in getattr(matrix, "figure_targets", [])}
    source_ids = _known_source_ids(source_catalog, source_status)
    policies = section_policy.by_section_id()
    table_lookup = _tables_by_id(comparison_tables)
    table_lookup.update(_included_tables_by_id(included))
    figure_lookup = _figures_by_id(map_manifest, included)
    issues: list[dict[str, Any]] = []

    if include_draft:
        issues.append(
            _qa_issue(
                "warning",
                "preview_export_qa_bypassed",
                "Internal preview export bypasses reviewed-export QA blocking; do not use it as reviewed content.",
            )
        )

    if not include_draft and review_gate.get("review_gate_status") != "passed":
        issues.append(
            _qa_issue(
                "error",
                "review_gate_not_passed",
                "Reviewed export requires a passed standard review gate before export QA can pass.",
            )
        )

    for item in included:
        item_id = str(item.get("id") or item.get("target_id") or "")
        item_type = str(item.get("type") or "")
        manual_material = item.get("manual_material", {}) if isinstance(item.get("manual_material"), dict) else {}
        manual_status = str(manual_material.get("material_status") or "")
        export_behavior = str(manual_material.get("export_behavior") or "")
        content_source = str(item.get("content_source") or "")
        if (
            item_type in {"report_section", "section_text"}
            and manual_status in {"manual_required", "restricted_reviewer_supplied_required"}
            and export_behavior == "body_replacement_when_reviewed"
            and content_source == "generated_content"
        ):
            issues.append(
                _qa_issue(
                    "error",
                    "manual_material_without_reviewer_content",
                    "Manual/restricted body content cannot export from generated placeholder text; reviewer-supplied edited or replacement content is required.",
                    item_id=item_id,
                )
            )

        _qa_check_refs(
            issues,
            item=item,
            item_id=item_id,
            table_ids=table_ids,
            figure_ids=figure_ids,
            source_ids=source_ids,
            policies=policies,
        )

        required_caveats = _required_caveats_for_export_item(item, policies)
        carried_caveats = set(_string_list(item.get("required_caveats", [])))
        missing_caveats = sorted(set(required_caveats) - carried_caveats)
        if item_type in {"report_section", "section_text"} and missing_caveats:
            issues.append(
                _qa_issue(
                    "error",
                    "required_caveat_missing",
                    "Included section is missing required caveat metadata: " + ", ".join(missing_caveats) + ".",
                    item_id=item_id,
                )
            )

        if item_type in {"comparison_table", "table"}:
            _qa_check_table_item(issues, item, item_id, table_lookup)
        if item_type in {"map_figure", "figure"}:
            _qa_check_figure_item(issues, project_dir, item, item_id, figure_lookup)

    for issue in validation_issues:
        if str(issue.get("severity") or "") == "error":
            issues.append(
                _qa_issue(
                    "error",
                    str(issue.get("code") or "export_validation_error"),
                    str(issue.get("message") or "Export validation reported an error."),
                    item_id=str(issue.get("item_id") or ""),
                    target_id=str(issue.get("target_id") or ""),
                )
            )
    for issue in _dict_list((final_verification or {}).get("issues", [])):
        if str(issue.get("severity") or "") == "error":
            issues.append(
                _qa_issue(
                    "error",
                    str(issue.get("code") or "final_verification_error"),
                    str(issue.get("message") or "Final verification reported an error."),
                )
            )
    if final_verification and final_verification.get("status") == "failed":
        issues.append(_qa_issue("error", "final_verification_failed", "Final verification status is failed."))

    if not compactness_budget:
        issues.append(_qa_issue("error", "compactness_budget_missing", "Export QA requires compactness budget metadata."))

    issues = _dedupe_issues(issues)
    blocking = [issue for issue in issues if str(issue.get("severity")) == "error"]
    warning = [issue for issue in issues if str(issue.get("severity")) == "warning"]
    return {
        "export_qa_status": "failed" if blocking else ("warning" if warning else "passed"),
        "preview_mode": include_draft,
        "checked_at": _utc_now(),
        "export_qa_issue_count": len(issues),
        "export_qa_blocking_error_count": len(blocking),
        "export_qa_warning_count": len(warning),
        "export_qa_issues": issues,
        "hard_block_policy": "reviewed_export_blocks_on_error",
        "override_supported": False,
        "override_policy": "deferred",
    }


def _qa_check_refs(
    issues: list[dict[str, Any]],
    *,
    item: dict[str, Any],
    item_id: str,
    table_ids: set[str],
    figure_ids: set[str],
    source_ids: set[str],
    policies: dict[str, Any],
) -> None:
    table_refs = _string_list(item.get("related_table_ids", []))
    if item.get("table_id"):
        table_refs.append(str(item.get("table_id")))
    for table_id in table_refs:
        if table_id and table_id not in table_ids:
            issues.append(_qa_issue("error", "unknown_table_ref", f"Included item references unknown table id '{table_id}'.", item_id=item_id, target_id=table_id))
    figure_refs = _string_list(item.get("related_figure_ids", []))
    if item.get("figure_id"):
        figure_refs.append(str(item.get("figure_id")))
    for figure_id in figure_refs:
        if figure_id and figure_id not in figure_ids and not _is_attachment_panel_figure_id(figure_id):
            issues.append(_qa_issue("error", "unknown_figure_ref", f"Included item references unknown figure id '{figure_id}'.", item_id=item_id, target_id=figure_id))
    for source_ref in _string_list(item.get("source_refs", [])):
        if source_ref and source_ref not in source_ids:
            issues.append(_qa_issue("error", "unknown_source_ref", f"Included item references unknown source id '{source_ref}'.", item_id=item_id, target_id=source_ref))

    policy = _policy_for_export_item(item, policies)
    if policy is None:
        return
    allowed_tables = set(_string_list(item.get("allowed_table_refs", [])) or list(policy.allowed_table_refs))
    allowed_figures = set(_string_list(item.get("allowed_figure_refs", [])) or list(policy.allowed_figure_refs))
    if allowed_tables:
        for table_id in _string_list(item.get("related_table_ids", [])):
            if table_id not in allowed_tables:
                issues.append(_qa_issue("error", "disallowed_table_ref", f"Included section references table id '{table_id}' outside its policy allowance.", item_id=item_id, target_id=table_id))
    if allowed_figures:
        for figure_id in _string_list(item.get("related_figure_ids", [])):
            if figure_id not in allowed_figures:
                issues.append(_qa_issue("error", "disallowed_figure_ref", f"Included section references figure id '{figure_id}' outside its policy allowance.", item_id=item_id, target_id=figure_id))


def _known_source_ids(source_catalog: Any, source_status: dict[str, Any]) -> set[str]:
    values = set(getattr(source_catalog, "sources", {}).keys())
    for status in _dict_list(source_status.get("statuses", [])):
        values.update(_string_list(status.get("source_ids", [])))
        values.update(_string_list(status.get("registered_source_ids", [])))
        for detail in _dict_list(status.get("source_details", [])):
            source_id = str(detail.get("source_id") or "").strip()
            if source_id:
                values.add(source_id)
    return values


def _is_attachment_panel_figure_id(value: str) -> bool:
    if not value.startswith("attachment-a-panel-"):
        return False
    suffix = value.removeprefix("attachment-a-panel-")
    return len(suffix) == 3 and suffix.isdigit()


def _qa_check_table_item(
    issues: list[dict[str, Any]],
    item: dict[str, Any],
    item_id: str,
    table_lookup: dict[str, dict[str, Any]],
) -> None:
    if bool(item.get("is_stub", False)):
        return
    columns = _string_list(item.get("columns", []))
    rows = _dict_list(item.get("rows_preview", []))
    table = table_lookup.get(str(item.get("table_id") or ""))
    if table:
        columns = columns or _string_list(table.get("columns", []))
        rows = rows or _dict_list(table.get("rows", []))
    required_columns = _string_list(item.get("required_columns", [])) or _string_list((table or {}).get("required_columns", []))
    for row_index, row in enumerate(rows, start=1):
        for column in required_columns:
            if not str(row.get(column, "")).strip():
                issues.append(_qa_issue("error", "blank_required_table_cell", f"Included table has a blank required cell in column '{column}' row {row_index}.", item_id=item_id, target_id=str(item.get("table_id") or "")))


def _qa_check_figure_item(
    issues: list[dict[str, Any]],
    project_dir: Path,
    item: dict[str, Any],
    item_id: str,
    figure_lookup: dict[str, dict[str, Any]],
) -> None:
    if bool(item.get("is_stub", False)):
        return
    figure_id = _item_figure_id(item)
    figure = figure_lookup.get(figure_id, {})
    caption = str(item.get("caption") or figure.get("caption") or "").strip()
    source_note = str(item.get("source_note") or figure.get("source_note") or "").strip()
    method_note = str(item.get("method_note") or figure.get("method_note") or "").strip()
    image_path = str(item.get("image_path") or figure.get("image_path") or "").strip()
    for field, value in (("caption", caption), ("source_note", source_note), ("method_note", method_note)):
        if not value:
            issues.append(_qa_issue("error", f"missing_figure_{field}", f"Included non-stub figure is missing {field}.", item_id=item_id, target_id=figure_id))
    if not image_path:
        if _figure_placeholder_allowed(source_note=source_note, method_note=method_note):
            issues.append(
                _qa_issue(
                    "warning",
                    "figure_placeholder_without_image",
                    "Included figure is reviewed as a placeholder/status item without an image path.",
                    item_id=item_id,
                    target_id=figure_id,
                )
            )
        else:
            issues.append(_qa_issue("error", "missing_figure_image_path", "Included non-stub figure is missing image_path.", item_id=item_id, target_id=figure_id))
    elif not _resolve_project_path(project_dir, image_path).exists():
        issues.append(_qa_issue("error", "missing_figure_image_file", f"Included non-stub figure image file is missing: {image_path}.", item_id=item_id, target_id=figure_id))


def _figure_placeholder_allowed(*, source_note: str, method_note: str) -> bool:
    text = f"{source_note} {method_note}".lower()
    return any(marker in text for marker in ("placeholder", "unavailable", "unsupported for automated figure rendering"))


def _required_caveats_for_export_item(item: dict[str, Any], policies: dict[str, Any]) -> list[str]:
    policy = _policy_for_export_item(item, policies)
    if policy:
        return list(policy.required_caveats)
    return _string_list(item.get("required_caveats", []))


def _policy_for_export_item(item: dict[str, Any], policies: dict[str, Any]) -> Any | None:
    target_id = str(item.get("target_id") or item.get("id") or "")
    if target_id in policies:
        return policies[target_id]
    provenance = item.get("provenance", {}) if isinstance(item.get("provenance"), dict) else {}
    deliverable_provenance = provenance.get("deliverable_item_provenance", {}) if isinstance(provenance.get("deliverable_item_provenance"), dict) else {}
    template_id = str(deliverable_provenance.get("template_target_id") or "")
    return policies.get(template_id) if template_id else None


def _qa_issue(
    severity: str,
    code: str,
    message: str,
    *,
    item_id: str = "",
    target_id: str = "",
) -> dict[str, str]:
    issue = _issue(severity, code, message)
    if item_id:
        issue["item_id"] = item_id
    if target_id:
        issue["target_id"] = target_id
    return issue


def _export_qa_error_message(details: dict[str, Any]) -> str:
    issues = _dict_list(details.get("export_qa_issues", []))
    blocking = [issue for issue in issues if str(issue.get("severity") or "") == "error"]
    examples = ", ".join(
        f"{issue.get('code')}{' on ' + str(issue.get('item_id')) if issue.get('item_id') else ''}"
        for issue in blocking[:5]
    )
    suffix = f" First issues: {examples}." if examples else ""
    return (
        "Reviewed export is blocked by export QA: "
        f"{details.get('export_qa_blocking_error_count', 0)} blocking issue(s)."
        f"{suffix} Use --include-draft only for an internal/pre-review preview."
    )


def _unresolved_required_sources(source_status: dict[str, Any]) -> list[dict[str, Any]]:
    unresolved: list[dict[str, Any]] = []
    for item in _dict_list(source_status.get("statuses", [])):
        if item.get("requirement") != "required" or str(item.get("status")) not in UNRESOLVED_REQUIRED_SOURCE_STATUSES:
            continue
        unresolved.append(
            {
                "category": item.get("category"),
                "status": item.get("status"),
                "source_ids": _string_list(item.get("source_ids", [])),
                "notes": item.get("notes", ""),
                "uncertainty_flags": _string_list(item.get("uncertainty_flags", [])),
            }
        )
    return unresolved


def _markdown_report(
    *,
    queue: dict[str, Any],
    included: list[dict[str, Any]],
    validation_issues: list[dict[str, Any]],
    include_draft: bool,
    data_lineage: dict[str, Any],
    comparison_tables: dict[str, Any] | None,
    map_manifest: dict[str, Any] | None,
    figure_assets: list[dict[str, Any]],
) -> str:
    lines = [
        f"# {queue.get('project_name', 'Environmental Constraints Report')}",
        "",
        "Environmental Constraints Report",
        "",
    ]
    if include_draft:
        lines.extend(
            [
                "> INTERNAL PREVIEW EXPORT: this file includes draft or unaccepted review queue items and is not ready for external use.",
                f"> Preview status: {len(included)} items are shown for inspection; {_queue_unreviewed_count(queue)} review queue item(s) still require reviewer action before reviewed export.",
                "",
            ]
        )
    else:
        lines.extend(
            [
                "> Reviewed-content export: this file includes only accepted, edited, or explicitly export-eligible reviewed items.",
                "",
            ]
        )
    if validation_issues:
        lines.append("## Export Caveats")
        lines.extend(f"- {issue.get('message') or issue.get('code')}" for issue in validation_issues)
        lines.append("")
    lines.extend(_markdown_data_lineage(data_lineage))
    if not included:
        lines.extend(["## No Exported Content", "", "No review queue items met the export criteria.", ""])
        return "\n".join(lines)

    table_lookup = _tables_by_id(comparison_tables)
    table_lookup.update(_included_tables_by_id(included))
    figure_lookup = _figures_by_id(map_manifest, included)
    included_table_ids = _included_table_ids(included)
    included_figure_ids = _included_figure_ids(included)
    planned_inline_table_ids = _inline_table_ids(included, included_table_ids, table_lookup)
    planned_inline_figure_ids = _inline_figure_ids(included, included_figure_ids, figure_lookup)
    rendered_table_ids: set[str] = set()
    rendered_figure_ids: set[str] = set()
    current_group = ""
    for item in included:
        group = str(item.get("export_group", "resource_sections"))
        if group != current_group:
            current_group = group
            if not _structural_heading_matches_group(item, group):
                lines.extend([f"## {EXPORT_GROUP_TITLES.get(group, group.replace('_', ' ').title())}", ""])
        if _standalone_rendered_inline(item, planned_inline_table_ids, planned_inline_figure_ids):
            continue
        lines.extend(
            _markdown_item(
                item,
                suppress_heading=_suppress_item_heading(item),
                table_lookup=table_lookup,
                figure_lookup=figure_lookup,
                included_table_ids=included_table_ids,
                included_figure_ids=included_figure_ids,
                rendered_table_ids=rendered_table_ids,
                rendered_figure_ids=rendered_figure_ids,
                all_included_items=included,
            )
        )
    return "\n".join(lines).rstrip() + "\n"


def _markdown_item(
    item: dict[str, Any],
    *,
    suppress_heading: bool = False,
    table_lookup: dict[str, dict[str, Any]] | None = None,
    figure_lookup: dict[str, dict[str, Any]] | None = None,
    included_table_ids: set[str] | None = None,
    included_figure_ids: set[str] | None = None,
    rendered_table_ids: set[str] | None = None,
    rendered_figure_ids: set[str] | None = None,
    all_included_items: list[dict[str, Any]] | None = None,
) -> list[str]:
    lines: list[str] = []
    heading = _item_heading_text(item)
    if not suppress_heading:
        heading_marker = "##" if item.get("type") == "structural_heading" else "###"
        lines.extend([f"{heading_marker} {heading}", ""])
    if item.get("type") == "structural_heading":
        return lines
    front_matter_list_kind = _front_matter_list_kind(item)
    content = _strip_duplicate_leading_heading(str(item.get("content", "")).strip(), heading, str(item.get("title", "")))
    if content and not front_matter_list_kind:
        lines.extend([content, ""])
    if item.get("type") in {"report_section", "section_text"} and not _is_front_matter_item(item):
        lines.extend(
            _markdown_inline_evidence(
                item,
                table_lookup=table_lookup or {},
                figure_lookup=figure_lookup or {},
                included_table_ids=included_table_ids or set(),
                included_figure_ids=included_figure_ids or set(),
                rendered_table_ids=rendered_table_ids if rendered_table_ids is not None else set(),
                rendered_figure_ids=rendered_figure_ids if rendered_figure_ids is not None else set(),
            )
        )
    if front_matter_list_kind:
        lines.extend(_markdown_front_matter_list(front_matter_list_kind, all_included_items or []))
    elif item.get("type") in {"report_section", "section_text"} and _is_front_matter_item(item):
        lines.extend(_markdown_front_matter_lists(all_included_items or []))
    if item.get("type") in {"map_figure", "figure"}:
        lines.extend(_markdown_embedded_figure(item))
    if item.get("type") in {"comparison_table", "table"}:
        details = []
        if item.get("row_count") is not None:
            details.append(f"{item['row_count']} row(s)")
        if details:
            lines.extend([f"Table summary: {', '.join(details)} are available for reviewer inspection.", ""])
    if not _is_front_matter_item(item):
        visual_slots = _string_list(item.get("visual_slots", []))
        related_figures = _string_list(item.get("related_figure_ids", []))
        missing_visuals = visual_slots if visual_slots and not related_figures else []
        if missing_visuals:
            lines.extend(["Visual needed:", *[f"- {slot}" for slot in missing_visuals], ""])
        table_slots = _string_list(item.get("table_slots", []))
        related_tables = _string_list(item.get("related_table_ids", []))
        missing_tables = table_slots if table_slots and not related_tables else []
        if missing_tables:
            lines.extend(["Table needed:", *[f"- {slot}" for slot in missing_tables], ""])
    if item.get("uncertainty_flags"):
        limitations = [_reviewer_flag_label(flag) for flag in _string_list(item.get("uncertainty_flags", []))]
        lines.extend(["Review limitations: " + ", ".join(limitations) + ".", ""])
    return lines


def _markdown_inline_evidence(
    item: dict[str, Any],
    *,
    table_lookup: dict[str, dict[str, Any]],
    figure_lookup: dict[str, dict[str, Any]],
    included_table_ids: set[str],
    included_figure_ids: set[str],
    rendered_table_ids: set[str],
    rendered_figure_ids: set[str],
) -> list[str]:
    lines: list[str] = []
    for table_id in _string_list(item.get("related_table_ids", [])):
        if table_id not in included_table_ids:
            lines.extend(["Table placeholder: a referenced support table was not included in this export.", ""])
            continue
        table = table_lookup.get(table_id)
        if table is None:
            lines.extend(["Table placeholder: a referenced support table was not available for export.", ""])
            continue
        if table_id in rendered_table_ids:
            lines.extend([f"Table reference: {_support_artifact_label(table, 'table_id', 'Table')} is provided above.", ""])
            continue
        lines.extend(_markdown_embedded_table(table))
        rendered_table_ids.add(table_id)
    for figure_id in _string_list(item.get("related_figure_ids", [])):
        if figure_id not in included_figure_ids:
            lines.extend(["Figure placeholder: a referenced support figure was not included in this export.", ""])
            continue
        figure = figure_lookup.get(figure_id)
        if figure is None:
            lines.extend(["Figure placeholder: a referenced support figure was not available for export.", ""])
            continue
        if figure_id in rendered_figure_ids:
            lines.extend([f"Figure reference: {_support_artifact_label(figure, 'figure_id', 'Figure')} is provided above.", ""])
            continue
        lines.extend(_markdown_embedded_figure(figure))
        rendered_figure_ids.add(figure_id)
    return lines


def _markdown_embedded_table(table: dict[str, Any]) -> list[str]:
    table_id = str(table.get("table_id") or "table")
    title = str(table.get("title") or table_id)
    rows = _dict_list(table.get("rows", []))
    columns = _string_list(table.get("columns", []))
    row_count = _optional_int(table.get("row_count"))
    total_rows = row_count if row_count is not None else len(rows)
    if not columns:
        columns = sorted({str(key) for row in rows for key in row})
    lines = [f"Table: {_support_artifact_label(table, 'table_id', 'Table')}", ""]
    if not columns:
        lines.extend(["Table placeholder: no columns were available for this table.", ""])
        return lines
    if not rows:
        lines.extend(["Table placeholder: this table currently has no rows.", ""])
        return lines
    rendered_rows = rows[: min(DOCX_TABLE_ROW_LIMIT, 10)]
    lines.append("| " + " | ".join(columns) + " |")
    lines.append("| " + " | ".join("---" for _ in columns) + " |")
    for row in rendered_rows:
        lines.append("| " + " | ".join(_markdown_cell(row.get(column)) for column in columns) + " |")
    if total_rows > len(rendered_rows):
        lines.append(f"Table preview limited to {len(rendered_rows)} of {total_rows} rows. Full table data remains in the table artifact.")
    lines.append("")
    return lines


def _markdown_embedded_figure(figure: dict[str, Any]) -> list[str]:
    figure_id = str(figure.get("figure_id") or "figure")
    title = str(figure.get("title") or figure_id)
    image_path = _markdown_figure_path(figure)
    if image_path:
        lines = [f"Figure: {_support_artifact_label(figure, 'figure_id', 'Figure')}", "", f"![{_markdown_alt_text(title)}]({image_path})", ""]
        caption = str(figure.get("caption") or "").strip()
        source_note = str(figure.get("source_note") or "").strip()
        method_note = str(figure.get("method_note") or "").strip()
        if caption:
            lines.extend([f"Caption: {caption}", ""])
        if source_note:
            lines.extend([f"Source note: {source_note}", ""])
        if method_note:
            lines.extend([f"Method note: {method_note}", ""])
        return lines
    return ["Figure placeholder: the selected figure image was not available for export.", ""]


def _support_artifact_label(record: dict[str, Any], id_field: str, kind: str) -> str:
    number_field = "table_number" if id_field == "table_id" else "figure_number"
    number = record.get(number_field)
    title = str(record.get("title") or "").strip()
    if number:
        prefix = f"{kind} {number}"
        return f"{prefix}. {title}" if title else prefix
    return title or f"the referenced {kind.lower()}"


def _reviewer_flag_label(flag: str) -> str:
    return flag.replace("_", " ").replace("-", " ")


def _queue_unreviewed_count(queue: dict[str, Any]) -> int:
    return sum(
        1
        for item in _dict_list(queue.get("items", []))
        if str(item.get("status") or "") not in TERMINAL_REVIEW_STATUSES | GATE_INCLUDED_NONTERMINAL_STATUSES
    )


def _markdown_figure_path(figure: dict[str, Any]) -> str:
    return str(figure.get("export_asset_path") or figure.get("export_image_path") or figure.get("image_path") or "").replace("\\", "/")


def _markdown_alt_text(value: str) -> str:
    return value.replace("[", "(").replace("]", ")")


def _markdown_front_matter_lists(included: list[dict[str, Any]]) -> list[str]:
    figures = [item for item in included if item.get("type") in {"map_figure", "figure"}]
    tables = [item for item in included if item.get("type") in {"comparison_table", "table"}]
    lines = ["List of Figures", ""]
    if figures:
        lines.extend(f"- {_front_matter_export_figure_label(item, index)}" for index, item in enumerate(figures, start=1))
    else:
        lines.append("- No figure items are included in this export.")
    lines.extend(["", "List of Tables", ""])
    if tables:
        lines.extend(f"- {_front_matter_export_table_label(item, index)}" for index, item in enumerate(tables, start=1))
    else:
        lines.append("- No table items are included in this export.")
    lines.extend(["", "List of Attachments", ""])
    attachments = _front_matter_attachment_items(included)
    if attachments:
        lines.extend(f"- {_attachment_list_label(item)}" for item in attachments)
    else:
        lines.append("- No attachment items are included in this export.")
    lines.append("")
    return lines


def _markdown_front_matter_list(kind: str, included: list[dict[str, Any]]) -> list[str]:
    if kind == "figures":
        figures = [item for item in included if item.get("type") in {"map_figure", "figure"}]
        if figures:
            return [*[f"- {_front_matter_export_figure_label(item, index)}" for index, item in enumerate(figures, start=1)], ""]
        return ["- No figure items are included in this export.", ""]
    if kind == "tables":
        tables = [item for item in included if item.get("type") in {"comparison_table", "table"}]
        if tables:
            return [*[f"- {_front_matter_export_table_label(item, index)}" for index, item in enumerate(tables, start=1)], ""]
        return ["- No table items are included in this export.", ""]
    if kind == "attachments":
        attachments = _front_matter_attachment_items(included)
        if attachments:
            return [*[f"- {_attachment_list_label(item)}" for item in attachments], ""]
        return ["- No attachment items are included in this export.", ""]
    return []


def _front_matter_export_figure_label(item: dict[str, Any], fallback_number: int) -> str:
    number = _front_matter_export_number(item, base=20_000, fallback=fallback_number)
    title = str(item.get("caption") or item.get("title") or "Untitled Figure").strip()
    return f"Figure {number}. {title}"


def _front_matter_export_table_label(item: dict[str, Any], fallback_number: int) -> str:
    number = _front_matter_export_number(item, base=10_000, fallback=fallback_number)
    title = str(item.get("title") or "Untitled Table").strip()
    return f"Table {number}. {title}"


def _front_matter_export_number(item: dict[str, Any], *, base: int, fallback: int) -> int:
    order = _optional_number(item.get("section_order"))
    if isinstance(order, (int, float)) and base < int(order) < base + 1000:
        return int(order) - base
    return fallback


def _markdown_data_lineage(data_lineage: dict[str, Any]) -> list[str]:
    lines = ["## Real Data Used", ""]
    real_records = [
        record
        for record in _dict_list(data_lineage.get("records", []))
        if record.get("lineage_type") in {"downloaded_public_source", "registered_local", "provided_in_input"}
    ]
    if real_records:
        for record in real_records:
            label = record.get("source_name") or record.get("source_id") or record.get("path") or "Real data source"
            details = []
            if record.get("status"):
                details.append(f"status `{record.get('status')}`")
            if record.get("feature_count") is not None:
                details.append(f"{record.get('feature_count')} feature(s)")
            if record.get("access_date"):
                details.append(f"accessed {record.get('access_date')}")
            lines.append(f"- {label}: {', '.join(details) if details else record.get('lineage_type')}.")
    else:
        lines.append("- No downloaded, provided-in-input, or registered local source layers were available.")
    lines.append("")
    stub_records = [
        record
        for record in _dict_list(data_lineage.get("records", []))
        if record.get("lineage_type") in {"manual_stub", "gated_stub", "missing_stub"}
    ]
    lines.extend(["## Stubs / Manual Review Needed", ""])
    if stub_records:
        for record in stub_records:
            label = record.get("category") or record.get("source_id") or "source category"
            lines.append(f"- {label}: {record.get('status', 'stub')} ({record.get('lineage_type')}).")
    else:
        lines.append("- No source stubs were recorded for this export.")
    lines.append("")
    return lines


def _markdown_package_contents(queue: dict[str, Any], *, figure_assets: list[dict[str, Any]] | None = None) -> list[str]:
    lines = ["## Generated Package Contents", ""]
    contents = _package_contents(queue, output_paths={}, figure_assets=figure_assets)
    for label, value in contents.items():
        if isinstance(value, dict):
            if not value:
                continue
            lines.append(f"- {label}:")
            for nested_label, nested_value in value.items():
                if nested_value:
                    lines.append(f"  - {nested_label}: `{nested_value}`")
        elif isinstance(value, list):
            if not value:
                continue
            lines.append(f"- {label}:")
            for item in value:
                if isinstance(item, dict):
                    item_label = item.get("title") or item.get("figure_id") or "asset"
                    item_path = item.get("export_asset_path") or item.get("export_image_path")
                    lines.append(f"  - {item_label}: `{item_path}`")
                else:
                    lines.append(f"  - `{item}`")
        elif value:
            lines.append(f"- {label}: `{value}`")
    lines.append("")
    return lines


def _write_docx_report(
    *,
    project_dir: Path,
    docx_path: Path,
    queue: dict[str, Any],
    included: list[dict[str, Any]],
    validation_issues: list[dict[str, Any]],
    include_draft: bool,
    comparison_tables: dict[str, Any] | None,
    map_manifest: dict[str, Any] | None,
    data_lineage: dict[str, Any],
    figure_assets: list[dict[str, Any]],
    output_paths: dict[str, str | None],
) -> None:
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:  # pragma: no cover - dependency is declared, this guards broken environments.
        raise ExportReportError("DOCX export requires the python-docx package to be installed.") from exc

    document = Document()
    _configure_docx_document(document, include_draft=include_draft)
    table_lookup = _tables_by_id(comparison_tables)
    table_lookup.update(_included_tables_by_id(included))
    figure_lookup = _figures_by_id(map_manifest, included)
    included_table_ids = _included_table_ids(included)
    included_figure_ids = _included_figure_ids(included)
    planned_inline_table_ids = _inline_table_ids(included, included_table_ids, table_lookup)
    planned_inline_figure_ids = _inline_figure_ids(included, included_figure_ids, figure_lookup)
    rendered_table_ids: set[str] = set()
    rendered_figure_ids: set[str] = set()

    _add_docx_title_page(document, project_dir, queue, include_draft=include_draft)

    if validation_issues:
        document.add_heading("Export Caveats", level=1)
        for issue in validation_issues:
            document.add_paragraph(str(issue.get("message") or issue.get("code")), style="List Bullet")

    _add_docx_data_lineage(document, data_lineage)

    if not included:
        document.add_heading("No Exported Content", level=1)
        document.add_paragraph("No review queue items met the export criteria.")
    else:
        current_group = ""
        first_group = True
        for item in included:
            group = str(item.get("export_group", "resource_sections"))
            if group != current_group:
                if not first_group:
                    document.add_page_break()
                first_group = False
                current_group = group
                if not _structural_heading_matches_group(item, group):
                    document.add_heading(EXPORT_GROUP_TITLES.get(group, group.replace("_", " ").title()), level=1)
            if _standalone_rendered_inline(item, planned_inline_table_ids, planned_inline_figure_ids):
                continue
            _add_docx_item(
                document=document,
                item=item,
                project_dir=project_dir,
                table_lookup=table_lookup,
                figure_lookup=figure_lookup,
                included_table_ids=included_table_ids,
                included_figure_ids=included_figure_ids,
                rendered_table_ids=rendered_table_ids,
                rendered_figure_ids=rendered_figure_ids,
                image_width=Inches(6.3),
                suppress_heading=_suppress_item_heading(item),
                all_included_items=included,
            )

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(docx_path)


def _add_docx_data_lineage(document: Any, data_lineage: dict[str, Any]) -> None:
    document.add_heading("Real Data Used", level=1)
    real_records = [
        record
        for record in _dict_list(data_lineage.get("records", []))
        if record.get("lineage_type") in {"downloaded_public_source", "registered_local", "provided_in_input"}
    ]
    if real_records:
        for record in real_records:
            label = record.get("source_name") or record.get("source_id") or record.get("path") or "Real data source"
            details = []
            if record.get("status"):
                details.append(f"status {record.get('status')}")
            if record.get("feature_count") is not None:
                details.append(f"{record.get('feature_count')} feature(s)")
            if record.get("access_date"):
                details.append(f"accessed {record.get('access_date')}")
            if record.get("checksum_sha256"):
                details.append(f"checksum {record.get('checksum_sha256')}")
            if record.get("source_limitations"):
                details.append(f"limitations: {record.get('source_limitations')}")
            document.add_paragraph(f"{label}: {'; '.join(details) if details else record.get('lineage_type')}.", style="List Bullet")
    else:
        document.add_paragraph("No downloaded, provided-in-input, or registered local source layers were available.", style="List Bullet")

    document.add_heading("Stubs / Manual Review Needed", level=1)
    stub_records = [
        record
        for record in _dict_list(data_lineage.get("records", []))
        if record.get("lineage_type") in {"manual_stub", "gated_stub", "missing_stub"}
    ]
    if stub_records:
        for record in stub_records:
            label = record.get("category") or record.get("source_id") or "source category"
            notes = str(record.get("notes") or "").strip()
            suffix = f" {notes}" if notes else ""
            document.add_paragraph(f"{label}: {record.get('status', 'stub')} ({record.get('lineage_type')}).{suffix}", style="List Bullet")
    else:
        document.add_paragraph("No source stubs were recorded for this export.", style="List Bullet")


def _configure_docx_document(document: Any, *, include_draft: bool) -> None:
    try:
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:  # pragma: no cover - guarded by caller.
        return

    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)
        if include_draft:
            section.header.paragraphs[0].text = "INTERNAL PREVIEW / NOT REVIEWED"
            section.footer.paragraphs[0].text = "Review Assist - Environmental Constraints Report - INTERNAL PREVIEW"
        else:
            section.footer.paragraphs[0].text = "Review Assist - Environmental Constraints Report"
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.16

    _set_docx_style(styles, "Title", font_name="Calibri Light", size=20, color=RGBColor(0x0F, 0x47, 0x61))
    _set_docx_style(styles, "Heading 1", font_name="Lato", size=20, color=RGBColor(0x0F, 0x47, 0x61), keep_with_next=True)
    _set_docx_style(styles, "Heading 2", font_name="Lato", size=14, color=RGBColor(0x0F, 0x47, 0x61), keep_with_next=True)
    _set_docx_style(styles, "Heading 3", font_name="Lato", color=RGBColor(0x0F, 0x47, 0x61), italic=True, keep_with_next=True)
    _set_docx_style(styles, "Heading 4", color=RGBColor(0x4C, 0x94, 0xD8), italic=True, keep_with_next=True)
    _set_docx_style(
        styles,
        "Caption",
        font_name="Calibri Light",
        size=11,
        color=RGBColor(0x0E, 0x28, 0x41),
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    try:
        styles["Attachment Title"]
    except KeyError:
        styles.add_style("Attachment Title", WD_STYLE_TYPE.PARAGRAPH)
    _set_docx_style(
        styles,
        "Attachment Title",
        color=RGBColor(0x0F, 0x47, 0x61),
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=True,
    )


def _set_docx_style(
    styles: Any,
    style_name: str,
    *,
    font_name: str | None = None,
    size: int | None = None,
    color: Any | None = None,
    italic: bool | None = None,
    bold: bool | None = None,
    alignment: Any | None = None,
    keep_with_next: bool = False,
) -> None:
    try:
        style = styles[style_name]
    except KeyError:
        return
    if font_name:
        style.font.name = font_name
    if size is not None:
        try:
            from docx.shared import Pt

            style.font.size = Pt(size)
        except ImportError:  # pragma: no cover - guarded by caller.
            pass
    if color is not None:
        style.font.color.rgb = color
    if italic is not None:
        style.font.italic = italic
    if bold is not None:
        style.font.bold = bold
    if alignment is not None:
        style.paragraph_format.alignment = alignment
    if keep_with_next:
        style.paragraph_format.keep_with_next = True


def _add_docx_title_page(document: Any, project_dir: Path, queue: dict[str, Any], *, include_draft: bool) -> None:
    try:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:  # pragma: no cover - guarded by caller.
        WD_ALIGN_PARAGRAPH = None

    metadata = _project_report_metadata(project_dir, queue)
    title = str(metadata.get("project_name") or "Environmental Constraints Report")
    title_paragraph = document.add_heading(title, level=0)
    if WD_ALIGN_PARAGRAPH is not None:
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    subtitle = document.add_paragraph("Environmental Constraints Report")
    if WD_ALIGN_PARAGRAPH is not None:
        subtitle.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph(f"Project ID: {metadata.get('project_id') or 'unknown'}")
    if metadata.get("description"):
        document.add_paragraph(str(metadata["description"]))
    if metadata.get("county_state_context"):
        document.add_paragraph(str(metadata["county_state_context"]))
    document.add_paragraph(f"Generated: {metadata['generated_at']}")
    if include_draft:
        notice = document.add_paragraph()
        run = notice.add_run("INTERNAL PREVIEW / NOT REVIEWED")
        run.bold = True
        document.add_paragraph(
            "This package includes draft or unaccepted review queue items. It is for internal preview only and is not ready for external use."
        )
        document.add_paragraph(
            f"Preview status: {_queue_unreviewed_count(queue)} review queue item(s) still require reviewer action before reviewed export."
        )
    else:
        document.add_paragraph(
            "Reviewed-content export: this package includes only accepted, edited, or explicitly export-eligible reviewed items."
        )
    document.add_paragraph(
        "This report presents objective desktop-screening constraints only and does not rank, select, reject, or recommend project features."
    )
    document.add_page_break()


def _project_report_metadata(project_dir: Path, queue: dict[str, Any]) -> dict[str, Any]:
    project_id = str(queue.get("project_id") or "")
    project_name = str(queue.get("project_name") or "Environmental Constraints Report")
    description = ""
    try:
        manifest = load_project_manifest(project_dir)
        project_id = manifest.project_id or project_id
        project_name = manifest.name or project_name
        description = manifest.description
    except ProjectManifestError:
        pass

    county_state_context = ""
    try:
        project_area = load_project_area(project_dir)
        counties = _string_list(project_area.get("county_names", []))
        if counties:
            county_state_context = f"County context: {', '.join(counties)}."
    except ProjectAreaError:
        pass

    return {
        "project_id": project_id,
        "project_name": project_name,
        "description": description,
        "county_state_context": county_state_context,
        "generated_at": _utc_now(),
    }


def _add_docx_item(
    *,
    document: Any,
    item: dict[str, Any],
    project_dir: Path,
    table_lookup: dict[str, dict[str, Any]],
    figure_lookup: dict[str, dict[str, Any]],
    included_table_ids: set[str],
    included_figure_ids: set[str],
    rendered_table_ids: set[str],
    rendered_figure_ids: set[str],
    image_width: Any,
    suppress_heading: bool = False,
    all_included_items: list[dict[str, Any]] | None = None,
) -> None:
    item_type = str(item.get("type", ""))
    heading = _item_heading_text(item)
    heading_level = _docx_heading_level(item, item_type)
    front_matter_list_kind = _front_matter_list_kind(item)
    if not suppress_heading:
        if item_type == "attachment":
            document.add_paragraph(heading, style="Attachment Title")
        else:
            document.add_heading(heading, level=heading_level)
    if item_type == "structural_heading":
        return

    content = _strip_duplicate_leading_heading(str(item.get("content", "")).strip(), heading, str(item.get("title", "")))
    if content and not front_matter_list_kind:
        _add_docx_content(document, content)
    elif item_type not in {"map_figure", "figure"} and not front_matter_list_kind:
        document.add_paragraph("No generated or reviewer-edited content was available for this item.")

    if item_type in {"comparison_table", "table"}:
        _add_docx_comparison_table(document, item, table_lookup)
    elif item_type in {"map_figure", "figure"}:
        _add_docx_map_figure(document, item, project_dir, image_width)
    else:
        if front_matter_list_kind:
            _add_docx_front_matter_list(document, front_matter_list_kind, all_included_items or [])
        elif item_type in {"report_section", "section_text"} and _is_front_matter_item(item):
            _add_docx_front_matter_lists(document, all_included_items or [])
        elif item_type in {"report_section", "section_text"}:
            _add_docx_inline_evidence(
                document=document,
                item=item,
                project_dir=project_dir,
                table_lookup=table_lookup,
                figure_lookup=figure_lookup,
                included_table_ids=included_table_ids,
                included_figure_ids=included_figure_ids,
                rendered_table_ids=rendered_table_ids,
                rendered_figure_ids=rendered_figure_ids,
                image_width=image_width,
            )
        if not _is_front_matter_item(item):
            _add_docx_missing_slots(document, item)

    if item.get("uncertainty_flags"):
        document.add_paragraph(
            "Review limitations: "
            + ", ".join(_reviewer_flag_label(flag) for flag in _string_list(item.get("uncertainty_flags", [])))
            + "."
        )


def _add_docx_inline_evidence(
    *,
    document: Any,
    item: dict[str, Any],
    project_dir: Path,
    table_lookup: dict[str, dict[str, Any]],
    figure_lookup: dict[str, dict[str, Any]],
    included_table_ids: set[str],
    included_figure_ids: set[str],
    rendered_table_ids: set[str],
    rendered_figure_ids: set[str],
    image_width: Any,
) -> None:
    for table_id in _string_list(item.get("related_table_ids", [])):
        if table_id not in included_table_ids:
            document.add_paragraph("Table placeholder: a referenced support table was not included in this export.")
            continue
        table = table_lookup.get(table_id)
        if table is None:
            document.add_paragraph("Table placeholder: a referenced support table was not available for export.")
            continue
        if table_id in rendered_table_ids:
            document.add_paragraph(f"Table reference: {_support_artifact_label(table, 'table_id', 'Table')} is provided above.")
            continue
        _add_docx_comparison_table(document, {"table_id": table_id}, table_lookup)
        rendered_table_ids.add(table_id)
    for figure_id in _string_list(item.get("related_figure_ids", [])):
        if figure_id not in included_figure_ids:
            document.add_paragraph("Figure placeholder: a referenced support figure was not included in this export.")
            continue
        figure = figure_lookup.get(figure_id)
        if figure is None:
            document.add_paragraph("Figure placeholder: a referenced support figure was not available for export.")
            continue
        if figure_id in rendered_figure_ids:
            document.add_paragraph(f"Figure reference: {_support_artifact_label(figure, 'figure_id', 'Figure')} is provided above.")
            continue
        document.add_paragraph(f"Figure: {_support_artifact_label(figure, 'figure_id', 'Figure')}", style="Caption")
        _add_docx_map_figure(document, figure, project_dir, image_width)
        rendered_figure_ids.add(figure_id)


def _add_docx_front_matter_lists(document: Any, included: list[dict[str, Any]]) -> None:
    figures = [item for item in included if item.get("type") in {"map_figure", "figure"}]
    tables = [item for item in included if item.get("type") in {"comparison_table", "table"}]
    document.add_heading("List of Figures", level=2)
    if figures:
        for index, item in enumerate(figures, start=1):
            document.add_paragraph(_front_matter_export_figure_label(item, index), style="List Bullet")
    else:
        document.add_paragraph("No figure items are included in this export.", style="List Bullet")
    document.add_heading("List of Tables", level=2)
    if tables:
        for index, item in enumerate(tables, start=1):
            document.add_paragraph(_front_matter_export_table_label(item, index), style="List Bullet")
    else:
        document.add_paragraph("No table items are included in this export.", style="List Bullet")
    document.add_heading("List of Attachments", level=2)
    attachments = _front_matter_attachment_items(included)
    if attachments:
        for item in attachments:
            document.add_paragraph(_attachment_list_label(item), style="List Bullet")
    else:
        document.add_paragraph("No attachment items are included in this export.", style="List Bullet")


def _add_docx_front_matter_list(document: Any, kind: str, included: list[dict[str, Any]]) -> None:
    if kind == "figures":
        figures = [item for item in included if item.get("type") in {"map_figure", "figure"}]
        if figures:
            for index, item in enumerate(figures, start=1):
                document.add_paragraph(_front_matter_export_figure_label(item, index), style="List Bullet")
        else:
            document.add_paragraph("No figure items are included in this export.", style="List Bullet")
    elif kind == "tables":
        tables = [item for item in included if item.get("type") in {"comparison_table", "table"}]
        if tables:
            for index, item in enumerate(tables, start=1):
                document.add_paragraph(_front_matter_export_table_label(item, index), style="List Bullet")
        else:
            document.add_paragraph("No table items are included in this export.", style="List Bullet")
    elif kind == "attachments":
        attachments = _front_matter_attachment_items(included)
        if attachments:
            for item in attachments:
                document.add_paragraph(_attachment_list_label(item), style="List Bullet")
        else:
            document.add_paragraph("No attachment items are included in this export.", style="List Bullet")


def _add_docx_content(document: Any, content: str) -> None:
    for raw_line in content.replace("\r\n", "\n").split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=4)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=3)
        elif line.startswith("# "):
            document.add_heading(line[2:].strip(), level=2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line)


def _add_docx_comparison_table(
    document: Any,
    item: dict[str, Any],
    table_lookup: dict[str, dict[str, Any]],
) -> None:
    table_id = str(item.get("table_id") or "")
    source_table = table_lookup.get(table_id)
    if source_table is None:
        document.add_paragraph("Table placeholder: a referenced support table was not available for export.")
        return

    document.add_paragraph(f"Table: {_support_artifact_label(source_table, 'table_id', 'Table')}", style="Caption")
    columns = _string_list(source_table.get("columns", []))
    rows = _dict_list(source_table.get("rows", []))
    row_count = _optional_int(source_table.get("row_count"))
    total_rows = row_count if row_count is not None else len(rows)
    if not columns:
        columns = sorted({str(key) for row in rows for key in row})
    if not columns:
        document.add_paragraph("Table placeholder: no columns were available for this table.")
        return
    if not rows:
        document.add_paragraph("Table placeholder: this table currently has no rows.")
        return

    rendered_rows = rows[:DOCX_TABLE_ROW_LIMIT]
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, column in enumerate(columns):
        header_cells[index].text = column
    for row in rendered_rows:
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            cells[index].text = _docx_cell_text(row.get(column))
    if total_rows > len(rendered_rows):
        document.add_paragraph(
            f"Table preview limited to {len(rendered_rows)} of {total_rows} rows. Full table data remains in the table artifact."
        )


def _add_docx_map_figure(document: Any, item: dict[str, Any], project_dir: Path, image_width: Any) -> None:
    image_value = item.get("export_image_path") or item.get("image_path")
    if not image_value:
        document.add_paragraph("Figure placeholder: no image path was recorded for this map figure.")
        return
    image_path = Path(str(image_value))
    if not image_path.is_absolute():
        image_path = project_dir / image_path
    if not image_path.exists():
        document.add_paragraph("Figure placeholder: the selected figure image was not available for export.")
        _add_docx_figure_notes(document, item)
        return
    try:
        document.add_picture(str(image_path), width=image_width)
    except Exception as exc:  # pragma: no cover - image backend errors vary by file.
        document.add_paragraph(f"Figure placeholder: the selected figure image could not be embedded: {exc}.")
        _add_docx_figure_notes(document, item)
        return
    _add_docx_figure_notes(document, item)


def _add_docx_figure_notes(document: Any, item: dict[str, Any]) -> None:
    caption = str(item.get("caption") or "").strip()
    source_note = str(item.get("source_note") or "").strip()
    method_note = str(item.get("method_note") or "").strip()
    if caption:
        document.add_paragraph(f"Caption: {caption}", style="Caption")
    if source_note:
        document.add_paragraph(f"Source note: {source_note}")
    if method_note:
        document.add_paragraph(f"Method note: {method_note}")


def _add_docx_missing_slots(document: Any, item: dict[str, Any]) -> None:
    visual_slots = _string_list(item.get("visual_slots", []))
    related_figures = _string_list(item.get("related_figure_ids", []))
    missing_visuals = visual_slots if visual_slots and not related_figures else []
    if missing_visuals:
        document.add_paragraph("Visual needed:")
        for slot in missing_visuals:
            document.add_paragraph(slot, style="List Bullet")

    table_slots = _string_list(item.get("table_slots", []))
    related_tables = _string_list(item.get("related_table_ids", []))
    missing_tables = table_slots if table_slots and not related_tables else []
    if missing_tables:
        document.add_paragraph("Table needed:")
        for slot in missing_tables:
            document.add_paragraph(slot, style="List Bullet")


def _add_docx_package_contents(
    document: Any,
    queue: dict[str, Any],
    output_paths: dict[str, str | None],
    *,
    figure_assets: list[dict[str, Any]] | None = None,
) -> None:
    document.add_heading("Generated Package Contents", level=1)
    contents = _package_contents(queue, output_paths=output_paths, figure_assets=figure_assets)
    for label, value in contents.items():
        if isinstance(value, dict):
            document.add_paragraph(f"{label}:")
            for nested_label, nested_value in value.items():
                if nested_value:
                    document.add_paragraph(f"{nested_label}: {nested_value}", style="List Bullet")
        elif isinstance(value, list):
            if value:
                document.add_paragraph(f"{label}:")
            for item in value:
                if isinstance(item, dict):
                    item_label = item.get("title") or item.get("figure_id") or "asset"
                    item_path = item.get("export_asset_path") or item.get("export_image_path")
                    document.add_paragraph(f"{item_label}: {item_path}", style="List Bullet")
                else:
                    document.add_paragraph(str(item), style="List Bullet")
        elif value:
            document.add_paragraph(f"{label}: {value}", style="List Bullet")


def _tables_by_id(comparison_tables: dict[str, Any] | None) -> dict[str, dict[str, Any]]:
    if comparison_tables is None:
        return {}
    return {
        str(table.get("table_id")): table
        for table in _dict_list(comparison_tables.get("tables", []))
        if table.get("table_id")
    }


def _included_tables_by_id(included: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    tables: dict[str, dict[str, Any]] = {}
    for item in included:
        if item.get("type") not in {"comparison_table", "table"} or not item.get("table_id"):
            continue
        table_id = str(item.get("table_id"))
        tables[table_id] = {
            "table_id": table_id,
            "title": item.get("title") or table_id,
            "columns": _string_list(item.get("columns", [])),
            "rows": _dict_list(item.get("rows_preview", [])),
            "required_columns": _string_list(item.get("required_columns", [])),
            "row_count": item.get("row_count"),
        }
    return tables


def _figures_by_id(map_manifest: dict[str, Any] | None, included: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    figures: dict[str, dict[str, Any]] = {}
    if map_manifest is not None:
        for figure in _dict_list(map_manifest.get("figures", [])):
            figure_id = str(figure.get("figure_id") or "").strip()
            if figure_id:
                figures[figure_id] = figure
    for item in included:
        if item.get("type") not in {"map_figure", "figure"}:
            continue
        figure_id = _item_figure_id(item)
        if not figure_id:
            continue
        figures[figure_id] = {
            **figures.get(figure_id, {}),
            "figure_id": figure_id,
            "title": item.get("title") or figures.get(figure_id, {}).get("title"),
            "image_path": item.get("image_path") or figures.get(figure_id, {}).get("image_path"),
            "source_refs": item.get("source_refs", []),
            "caption": item.get("caption") or figures.get(figure_id, {}).get("caption"),
            "source_note": item.get("source_note") or figures.get(figure_id, {}).get("source_note"),
            "method_note": item.get("method_note") or figures.get(figure_id, {}).get("method_note"),
            "map_elements": item.get("map_elements") or figures.get(figure_id, {}).get("map_elements", []),
            "figure_group": item.get("figure_group") or figures.get(figure_id, {}).get("figure_group"),
            "related_resource_categories": item.get("related_resource_categories")
            or figures.get(figure_id, {}).get("related_resource_categories", []),
            "export_image_path": item.get("export_image_path") or figures.get(figure_id, {}).get("export_image_path"),
            "export_asset_path": item.get("export_asset_path") or figures.get(figure_id, {}).get("export_asset_path"),
            **_selected_figure_version_fields(item),
        }
    return figures


def _included_table_ids(included: list[dict[str, Any]]) -> set[str]:
    return {str(item.get("table_id")) for item in included if item.get("type") in {"comparison_table", "table"} and item.get("table_id")}


def _included_figure_ids(included: list[dict[str, Any]]) -> set[str]:
    return {_item_figure_id(item) for item in included if item.get("type") in {"map_figure", "figure"} and _item_figure_id(item)}


def _included_attachment_ids(included: list[dict[str, Any]]) -> set[str]:
    values: set[str] = set()
    for item in included:
        if item.get("attachment_id"):
            values.add(str(item["attachment_id"]))
        for attachment_id in _string_list(item.get("attachment_refs", [])):
            values.add(attachment_id)
        if item.get("type") == "attachment" and item.get("id"):
            values.add(str(item["id"]))
    return values


def _final_verification_summary(
    *,
    included: list[dict[str, Any]],
    review_gate: dict[str, Any],
    compactness_budget: dict[str, Any],
    output_formats: list[str],
    markdown_path: Path | None,
    docx_path: Path | None,
) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    gate_status = str(review_gate.get("review_gate_status") or "")
    if gate_status not in {"passed", "preview_bypassed"}:
        issues.append(_issue("error", "final_review_gate_not_ready", "Final verification requires a passed review gate or explicit preview bypass."))
    if not compactness_budget:
        issues.append(_issue("error", "compactness_budget_missing", "Final verification requires the export compactness budget."))

    expected_count = review_gate.get("expected_deliverable_item_count")
    actual_count = review_gate.get("actual_deliverable_item_count")
    if expected_count is not None and actual_count is not None and expected_count != actual_count:
        issues.append(
            _issue(
                "error",
                "deliverable_item_count_mismatch",
                "Expected deliverable item count does not match the actual bounded review queue item count.",
            )
        )

    raw_items = [item for item in included if str(item.get("type") or "") in RAW_LEGACY_EXPORT_TYPES]
    if raw_items:
        issues.append(
            _issue(
                "error",
                "raw_legacy_items_in_export",
                "Raw legacy/audit review item types are present in the export candidate; standard deliverables must use bounded deliverable items.",
            )
        )

    max_table_preview_rows = 0
    for item in included:
        item_type = str(item.get("type") or "")
        if item_type == "table":
            if not item.get("table_id"):
                issues.append(_issue("error", "included_table_ref_missing", "Included table review item is missing table_id."))
            rows_preview = _dict_list(item.get("rows_preview", []))
            max_table_preview_rows = max(max_table_preview_rows, len(rows_preview))
            if len(rows_preview) > DELIVERABLE_TABLE_BODY_PREVIEW_LIMIT:
                issues.append(
                    _issue(
                        "error",
                        "table_preview_limit_exceeded",
                        f"Included table '{item.get('id')}' carries {len(rows_preview)} preview rows; body previews are capped at {DELIVERABLE_TABLE_BODY_PREVIEW_LIMIT}.",
                    )
                )
        elif item_type == "figure" and not _item_figure_id(item):
            issues.append(_issue("error", "included_figure_ref_missing", "Included figure review item is missing figure_id."))
        elif item_type == "attachment" and not item.get("attachment_id") and not _string_list(item.get("attachment_refs", [])):
            issues.append(_issue("error", "included_attachment_ref_missing", "Included attachment review item is missing attachment_id or attachment_refs."))

        content_length = len(str(item.get("content") or ""))
        if content_length > EXPORT_BODY_CONTENT_WARNING_CHAR_LIMIT:
            issues.append(
                _issue(
                    "warning",
                    "export_body_content_over_budget",
                    f"Included item '{item.get('id')}' content is {content_length} characters; compact report sections should stay under {EXPORT_BODY_CONTENT_WARNING_CHAR_LIMIT} characters unless reviewed.",
                )
            )

    if markdown_path is not None and not markdown_path.exists():
        issues.append(_issue("error", "markdown_export_missing", f"Markdown export was requested but not found at {markdown_path}."))

    docx_readable: bool | None = None
    if "docx" in output_formats:
        docx_readable = _docx_readability_check(docx_path, issues)

    severities = {str(issue.get("severity", "")) for issue in issues}
    status = "failed" if "error" in severities else ("warning" if "warning" in severities else "passed")
    return {
        "status": status,
        "checked_at": _utc_now(),
        "review_gate_status": gate_status,
        "compactness_budget_present": bool(compactness_budget),
        "expected_deliverable_item_count": expected_count,
        "actual_deliverable_item_count": actual_count,
        "included_item_count": len(included),
        "included_section_count": compactness_budget.get("included_section_count", 0) if compactness_budget else 0,
        "included_table_count": compactness_budget.get("included_table_count", 0) if compactness_budget else 0,
        "included_figure_count": compactness_budget.get("included_figure_count", 0) if compactness_budget else 0,
        "included_attachment_count": compactness_budget.get("included_attachment_count", 0) if compactness_budget else 0,
        "raw_legacy_item_count": len(raw_items),
        "raw_legacy_item_ids": [str(item.get("id") or "") for item in raw_items],
        "max_table_preview_rows": max_table_preview_rows,
        "table_preview_limit": DELIVERABLE_TABLE_BODY_PREVIEW_LIMIT,
        "docx_readable": docx_readable,
        "issue_count": len(issues),
        "issues": _dedupe_issues(issues),
    }


def _docx_readability_check(docx_path: Path | None, issues: list[dict[str, Any]]) -> bool:
    if docx_path is None:
        issues.append(_issue("error", "docx_export_path_missing", "DOCX export was requested but no DOCX path was provided."))
        return False
    if not docx_path.exists():
        issues.append(_issue("error", "docx_export_missing", f"DOCX export was requested but not found at {docx_path}."))
        return False
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - guarded by caller.
        issues.append(_issue("error", "docx_dependency_unavailable", f"DOCX readability check could not import python-docx: {exc}."))
        return False
    try:
        document = Document(docx_path)
    except Exception as exc:  # pragma: no cover - corrupt DOCX failures vary by library version.
        issues.append(_issue("error", "docx_export_unreadable", f"DOCX export could not be opened for final verification: {exc}."))
        return False
    if not document.paragraphs:
        issues.append(_issue("error", "docx_export_empty", "DOCX export opened but contained no paragraphs."))
        return False
    return True


def _stub_item_count(items: list[dict[str, Any]]) -> int:
    count = 0
    for item in items:
        flags = set(_string_list(item.get("uncertainty_flags", [])))
        if "deliverable_item_stub" in flags or _generated_content(item) == REQUIRED_STUB_TEXT:
            count += 1
    return count


def _compactness_budget(
    *,
    included: list[dict[str, Any]],
    review_gate: dict[str, Any],
    mvp_quality: dict[str, Any],
) -> dict[str, Any]:
    return {
        "included_section_count": int(mvp_quality.get("included_section_count", 0)),
        "included_table_count": int(mvp_quality.get("included_table_count", 0)),
        "included_figure_count": int(mvp_quality.get("included_figure_count", 0)),
        "included_attachment_count": len({str(item.get("attachment_id")) for item in included if item.get("attachment_id")}),
        "stub_item_count": int(review_gate.get("stub_item_count", 0)),
        "rendered_table_preview_row_count": _rendered_table_preview_row_count(included),
        "approximate_body_character_count": sum(len(str(item.get("content") or "")) for item in included),
        "review_gate_status": review_gate.get("review_gate_status"),
        "expected_deliverable_item_count": review_gate.get("expected_deliverable_item_count"),
        "actual_deliverable_item_count": review_gate.get("actual_deliverable_item_count"),
        "included_item_count": len(included),
    }


def _rendered_table_preview_row_count(included: list[dict[str, Any]]) -> int:
    total = 0
    for item in included:
        if item.get("type") not in {"comparison_table", "table"}:
            continue
        rows_preview = _dict_list(item.get("rows_preview", []))
        if not rows_preview:
            assumptions = item.get("assumptions", {}) if isinstance(item.get("assumptions"), dict) else {}
            rows_preview = _dict_list(assumptions.get("rows_preview", []))
        if rows_preview:
            total += len(rows_preview[:DOCX_TABLE_ROW_LIMIT])
    return total


def _export_map_path(item: dict[str, Any]) -> str:
    return str(item.get("export_asset_path") or item.get("export_image_path") or item.get("image_path") or "")


def _item_figure_id(item: dict[str, Any]) -> str:
    figure_id = str(item.get("figure_id") or "").strip()
    if figure_id:
        return figure_id
    provenance = item.get("provenance", {}) if isinstance(item.get("provenance"), dict) else {}
    figure_id = str(provenance.get("figure_id") or "").strip()
    if figure_id:
        return figure_id
    item_id = str(item.get("id") or "")
    if item_id.startswith("map-figure-"):
        return item_id.removeprefix("map-figure-")
    return ""


def _standalone_rendered_inline(
    item: dict[str, Any],
    planned_inline_table_ids: set[str],
    planned_inline_figure_ids: set[str],
) -> bool:
    if item.get("type") in {"comparison_table", "table"} and str(item.get("table_id") or "") in planned_inline_table_ids:
        return True
    if item.get("type") in {"map_figure", "figure"} and _item_figure_id(item) in planned_inline_figure_ids:
        return True
    return False


def _suppress_item_heading(item: dict[str, Any]) -> bool:
    if item.get("type") not in {"report_section", "section_text"}:
        return False
    group_title = EXPORT_GROUP_TITLES.get(str(item.get("export_group", "")), "")
    return _normalized_heading(str(item.get("title", ""))) == _normalized_heading(group_title)


def _structural_heading_matches_group(item: dict[str, Any], group: str) -> bool:
    if item.get("type") != "structural_heading":
        return False
    group_title = EXPORT_GROUP_TITLES.get(group, group.replace("_", " ").title())
    return _normalized_heading(str(item.get("title", ""))) == _normalized_heading(group_title)


def _is_front_matter_item(item: dict[str, Any]) -> bool:
    return item.get("type") in {"report_section", "section_text", "front_matter"} and str(item.get("export_group")) == "front_matter"


def _front_matter_list_kind(item: dict[str, Any]) -> str:
    target_id = str(item.get("target_id") or item.get("id") or "").strip()
    return {
        "list-of-figures": "figures",
        "list-of-tables": "tables",
        "list-of-attachments": "attachments",
    }.get(target_id, "")


def _front_matter_attachment_items(included: list[dict[str, Any]]) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in included:
        section_number = str(item.get("section_number") or "").strip()
        is_attachment_section = (
            str(item.get("export_group") or "") == "attachments"
            and section_number.lower().startswith("attachment")
            and item.get("type") in {"report_section", "section_text", "attachment"}
        )
        if item.get("type") != "attachment" and not is_attachment_section:
            continue
        item_id = str(item.get("id") or item.get("target_id") or "")
        if not section_number and item.get("attachment_id"):
            continue
        key = section_number or item_id
        if key in seen:
            continue
        seen.add(key)
        items.append(item)
    return items


def _attachment_list_label(item: dict[str, Any]) -> str:
    section_number = str(item.get("section_number") or "").strip()
    title = str(item.get("title") or item.get("attachment_id") or item.get("id") or "Attachment").strip()
    suffix = "." if not title.endswith(".") else ""
    if section_number:
        return f"{section_number}: {title}{suffix}"
    return f"{title}{suffix}"


def _item_heading_text(item: dict[str, Any]) -> str:
    title = str(item.get("title") or "Untitled Item").strip()
    section_number = str(item.get("section_number") or "").strip()
    if section_number and not _normalized_heading(title).startswith(_normalized_heading(section_number)):
        return f"{section_number} {title}".strip()
    return title or "Untitled Item"


def _docx_heading_level(item: dict[str, Any], item_type: str) -> int:
    heading_level = _optional_int(item.get("heading_level"))
    if heading_level is not None and 1 <= heading_level <= 4:
        return heading_level
    if item_type in {"front_matter", "attachment"}:
        return 2
    if item_type in {"report_section", "section_text"}:
        return 2
    return 3


def _strip_duplicate_leading_heading(content: str, *titles: str) -> str:
    lines = content.splitlines()
    while lines and not lines[0].strip():
        lines.pop(0)
    if not lines:
        return ""
    title_keys = {_normalized_heading(title) for title in titles if title}
    while lines:
        first = lines[0].strip()
        if not first.startswith("#"):
            break
        heading = first.lstrip("#").strip()
        if _normalized_heading(heading) not in title_keys:
            break
        lines.pop(0)
        while lines and not lines[0].strip():
            lines.pop(0)
    return "\n".join(lines).strip()


def _normalized_heading(value: str) -> str:
    return " ".join(value.strip().lower().split())


def _package_contents(
    queue: dict[str, Any],
    *,
    output_paths: dict[str, str | None],
    figure_assets: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    upstream = queue.get("upstream_artifacts", {}) if isinstance(queue.get("upstream_artifacts"), dict) else {}
    return {
        "project_dir": queue.get("project_dir"),
        "review_queue": queue.get("output_path"),
        "source_inventory": upstream.get("source_inventory_path"),
        "source_status": upstream.get("source_status_path"),
        "constraint_results": upstream.get("constraint_results_path"),
        "comparison_tables": upstream.get("comparison_tables_path"),
        "map_manifest": upstream.get("map_manifest_path"),
        "report_sections": upstream.get("report_sections_path"),
        "evidence_package": upstream.get("evidence_package_path"),
        "exports": {key: value for key, value in output_paths.items() if value},
        "figure_assets": figure_assets or [],
    }


def _evidence_package_path(queue: dict[str, Any]) -> Any:
    upstream = queue.get("upstream_artifacts", {}) if isinstance(queue.get("upstream_artifacts"), dict) else {}
    return upstream.get("evidence_package_path")


def _gpt_drafting_summary(items: list[dict[str, Any]]) -> dict[str, Any]:
    section_items = [item for item in items if item.get("type") in {"report_section", "section_text"}]
    gpt_items = []
    models = set()
    rejected = 0
    for item in section_items:
        section_provenance = _item_drafting_provenance(item)
        if section_provenance.get("draft_provider") != "openai_responses":
            continue
        gpt_items.append(item)
        if section_provenance.get("gpt_model"):
            models.add(str(section_provenance.get("gpt_model")))
        elif section_provenance.get("model"):
            models.add(str(section_provenance.get("model")))
        if section_provenance.get("gpt_output_accepted") is False:
            rejected += 1
    return {
        "enabled": bool(gpt_items),
        "section_count": len(gpt_items),
        "accepted_section_count": len(gpt_items) - rejected,
        "rejected_section_count": rejected,
        "models": sorted(models),
    }


def _item_drafting_provenance(item: dict[str, Any]) -> dict[str, Any]:
    provenance = item.get("provenance", {}) if isinstance(item.get("provenance"), dict) else {}
    interpretive = provenance.get("gpt_interpretive_assist", {}) if isinstance(provenance.get("gpt_interpretive_assist"), dict) else {}
    if interpretive:
        return interpretive
    section_provenance = provenance.get("section_provenance", {}) if isinstance(provenance.get("section_provenance"), dict) else {}
    if section_provenance:
        return section_provenance
    deliverable_provenance = provenance.get("deliverable_item_provenance", {}) if isinstance(provenance.get("deliverable_item_provenance"), dict) else {}
    drafting = deliverable_provenance.get("drafting", {}) if isinstance(deliverable_provenance.get("drafting"), dict) else {}
    if drafting:
        return drafting
    return provenance.get("drafting", {}) if isinstance(provenance.get("drafting"), dict) else {}


def _docx_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        text = json.dumps(value, sort_keys=True)
    else:
        text = str(value)
    return text if len(text) <= 250 else text[:247] + "..."


def _markdown_cell(value: Any) -> str:
    text = _docx_cell_text(value).replace("\n", " ")
    return text.replace("|", "\\|")


def _default_export_group(item: dict[str, Any]) -> str:
    item_type = str(item.get("type", ""))
    if item_type == "report_section":
        assumptions = item.get("assumptions", {}) if isinstance(item.get("assumptions"), dict) else {}
        value = assumptions.get("export_group")
        if isinstance(value, str) and value.strip():
            return value
    if item_type in {"map_figure", "comparison_table", "draft_finding", "spatial_relationship", "no_mapped_relationships"}:
        return "constraints_inventory"
    if item_type in {"missing_data_placeholder", "source_inventory_note", "source_status_note"}:
        return "methodology"
    return "attachments"


def _mvp_quality_summary(
    *,
    included: list[dict[str, Any]],
    data_lineage: dict[str, Any],
    unresolved_required_sources: list[dict[str, Any]],
    validation_issues: list[dict[str, Any]],
    comparison_tables: dict[str, Any] | None,
    map_manifest: dict[str, Any] | None,
    figure_assets: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    included_table_ids = _included_table_ids(included)
    included_figure_ids = _included_figure_ids(included)
    table_lookup = _tables_by_id(comparison_tables)
    table_lookup.update(_included_tables_by_id(included))
    figure_lookup = _figures_by_id(map_manifest, included)
    inline_table_ids = _inline_table_ids(included, included_table_ids, table_lookup)
    inline_figure_ids = _inline_figure_ids(included, included_figure_ids, figure_lookup)
    placeholder_count = _export_placeholder_count(
        included,
        included_table_ids=included_table_ids,
        included_figure_ids=included_figure_ids,
        table_lookup=table_lookup,
        figure_lookup=figure_lookup,
    )
    resource_placeholder_sections = _placeholder_resource_sections(included)
    gpt_summary = _gpt_drafting_summary(included)
    return {
        "real_source_count": int(data_lineage.get("real_source_count", 0)),
        "source_backed_constraint_count": int(data_lineage.get("source_backed_constraint_count", 0)),
        "included_section_count": sum(1 for item in included if item.get("type") in {"report_section", "section_text"}),
        "included_table_count": len(included_table_ids),
        "included_figure_count": len(included_figure_ids),
        "copied_figure_asset_count": len(figure_assets or []),
        "copied_figure_asset_ids": sorted(str(asset.get("figure_id")) for asset in (figure_assets or []) if asset.get("figure_id")),
        "inline_rendered_table_count": len(inline_table_ids),
        "inline_rendered_figure_count": len(inline_figure_ids),
        "inline_rendered_table_ids": sorted(inline_table_ids),
        "inline_rendered_figure_ids": sorted(inline_figure_ids),
        "placeholder_count": placeholder_count,
        "placeholder_resource_section_count": len(resource_placeholder_sections),
        "placeholder_resource_section_ids": sorted(resource_placeholder_sections),
        "unresolved_required_source_count": len(unresolved_required_sources),
        "unresolved_required_source_categories": [str(item.get("category")) for item in unresolved_required_sources],
        "gpt_section_count": int(gpt_summary.get("section_count", 0)),
        "gpt_accepted_section_count": int(gpt_summary.get("accepted_section_count", 0)),
        "gpt_rejected_section_count": int(gpt_summary.get("rejected_section_count", 0)),
        "validation_warning_count": sum(1 for issue in validation_issues if str(issue.get("severity", "warning")) == "warning"),
        "missing_figure_asset_warning_count": sum(
            1
            for issue in validation_issues
            if str(issue.get("code", "")) in {"missing_export_figure_source", "missing_export_figure_asset", "export_figure_asset_copy_failed"}
        ),
    }


def _mvp_quality_validation_issues(mvp_quality: dict[str, Any], *, include_draft: bool) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    real_source_count = int(mvp_quality.get("real_source_count", 0))
    if real_source_count > 0 and int(mvp_quality.get("source_backed_constraint_count", 0)) == 0:
        issues.append(_issue("warning", "mvp_no_source_backed_constraints", "MVP preview has real source layers but no source-backed constraints."))
    if include_draft and real_source_count > 0 and int(mvp_quality.get("inline_rendered_figure_count", 0)) == 0:
        issues.append(_issue("warning", "mvp_no_inline_figures", "MVP preview has real source layers but no figures rendered inline in report sections."))
    if include_draft and real_source_count > 0 and int(mvp_quality.get("inline_rendered_table_count", 0)) == 0:
        issues.append(_issue("warning", "mvp_no_inline_tables", "MVP preview has real source layers but no tables rendered inline in report sections."))
    if int(mvp_quality.get("included_figure_count", 0)) > int(mvp_quality.get("copied_figure_asset_count", 0)):
        issues.append(_issue("warning", "mvp_missing_export_figure_assets", "One or more included map figures were not copied into the export figure assets folder."))
    if int(mvp_quality.get("placeholder_resource_section_count", 0)) > 5:
        issues.append(_issue("warning", "mvp_many_placeholder_sections", "MVP preview still has many placeholder-only resource sections."))
    return issues


def _inline_table_ids(
    included: list[dict[str, Any]],
    included_table_ids: set[str],
    table_lookup: dict[str, dict[str, Any]],
) -> set[str]:
    values: set[str] = set()
    for item in included:
        if item.get("type") not in {"report_section", "section_text"}:
            continue
        for table_id in _string_list(item.get("related_table_ids", [])):
            if table_id in included_table_ids:
                values.add(table_id)
    return values


def _inline_figure_ids(
    included: list[dict[str, Any]],
    included_figure_ids: set[str],
    figure_lookup: dict[str, dict[str, Any]],
) -> set[str]:
    values: set[str] = set()
    for item in included:
        if item.get("type") not in {"report_section", "section_text"}:
            continue
        for figure_id in _string_list(item.get("related_figure_ids", [])):
            if figure_id in included_figure_ids and figure_id in figure_lookup:
                values.add(figure_id)
    return values


def _export_placeholder_count(
    included: list[dict[str, Any]],
    *,
    included_table_ids: set[str],
    included_figure_ids: set[str],
    table_lookup: dict[str, dict[str, Any]],
    figure_lookup: dict[str, dict[str, Any]],
) -> int:
    count = 0
    for item in included:
        related_tables = set(_string_list(item.get("related_table_ids", [])))
        related_figures = set(_string_list(item.get("related_figure_ids", [])))
        for table_id in related_tables:
            if table_id not in included_table_ids or table_id not in table_lookup:
                count += 1
        for figure_id in related_figures:
            if figure_id not in included_figure_ids or figure_id not in figure_lookup:
                count += 1
        if item.get("type") in {"report_section", "section_text"}:
            visual_slots = set(_string_list(item.get("visual_slots", [])))
            table_slots = set(_string_list(item.get("table_slots", [])))
            if visual_slots and not related_figures:
                count += len(visual_slots)
            if table_slots and not related_tables:
                count += len(table_slots)
    return count


def _placeholder_resource_sections(included: list[dict[str, Any]]) -> set[str]:
    section_ids: set[str] = set()
    for item in included:
        if item.get("type") not in {"report_section", "section_text"} or item.get("export_group") != "resource_sections":
            continue
        if item.get("source_refs") or item.get("related_table_ids") or item.get("related_figure_ids"):
            continue
        if item.get("visual_slots") or item.get("table_slots"):
            section_ids.add(str(item.get("id") or item.get("title") or "resource_section"))
    return section_ids


def _issue(severity: str, code: str, message: str) -> dict[str, str]:
    return {
        "severity": severity,
        "code": code,
        "message": message,
    }


def _dedupe_issues(issues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    deduped: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, str, str]] = set()
    for issue in issues:
        key = (
            str(issue.get("severity", "")),
            str(issue.get("code", "")),
            str(issue.get("message", "")),
            str(issue.get("location", "")),
            str(issue.get("item_id", "")),
        )
        if key in seen:
            continue
        seen.add(key)
        deduped.append(issue)
    return deduped


def _dict_list(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]


def _string_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item) for item in value if str(item).strip()]


def _optional_int(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _optional_number(value: Any) -> int | float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if number.is_integer() else number


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _slug(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", value.strip().lower()).strip("-")
    return slug or "item"
